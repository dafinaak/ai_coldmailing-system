#!/usr/bin/env python3
"""One Excel per ten zones: all contacts of 40-49 in one list, 50-59 in the
next, and so on (Dafina, 29.09.2026: "me i bo ne nje list te vetme te
gjitha kontaktet ... prej 40 deri 49 ne nje excel pastaj per 50 deri 59
tjetrin e kshtu me radhe" - "boni kejt te nje sheet mos i ndaj hiq"): one
sheet, the zones under each other, a last column "Zona".

Usage:
    .venv/bin/python werkzeuge/zonen-sammelliste.py          40-49, 50-59, 60-69
    .venv/bin/python werkzeuge/zonen-sammelliste.py 40       only 40-49

Run it again after every new zone: it always takes the newest final list
of each zone, so the merged file grows with the zones. The merging and its
checks live in pipeline/zone_lists.py.

Writes IT-Liste-Emails-Zonat-40-49-<time>.xlsx next to the zone lists and
changes none of them. No email is sent, Instantly is not touched.
"""
import sys
from datetime import datetime
from pathlib import Path

import openpyxl

PROJECT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT))

from pipeline.zone_lists import merge_zone_lists, newest_per_zone  # noqa: E402

# What each zone list holds - shown so the merged count can be checked.
SHEETS = ("Versandfertig", "Zur Kontrolle", "Bezahlt, nicht versandfähig")


def main():
    starts = [int(a) for a in sys.argv[1:]] or [40, 50, 60]
    for start in starts:
        files = newest_per_zone(PROJECT, start)
        print(f"=== Zonat {start}-{start + 9} ===")
        if not files:
            print("    ende asnje zone me liste perfundimtare - s'krijohet asgje")
            continue
        for zone, path in files.items():
            book = openpyxl.load_workbook(path, read_only=True)
            counts = ", ".join(f"{name}: {book[name].max_row - 1}"
                               for name in SHEETS if name in book.sheetnames)
            print(f"    zona {zone}: {path.name} ({counts})")
        stamp = datetime.now().strftime("%Y%m%d-%H%M")
        target = PROJECT / f"IT-Liste-Emails-Zonat-{start}-{start + 9}-{stamp}.xlsx"
        try:
            merge_zone_lists(files, target)
        except ValueError as error:
            sys.exit(f"    NDALIM: {error}")
        book = openpyxl.load_workbook(target, read_only=True)
        totals = ", ".join(f"{sheet.title}: {sheet.max_row - 1}"
                           for sheet in book.worksheets)
        print(f"    -> {target.name} ({totals})")


if __name__ == "__main__":
    main()
