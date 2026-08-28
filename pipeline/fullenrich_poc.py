"""FullEnrich-only Anreicherung - Testlauf (Auftrag Dafina, 24.08.2026).

Frage: schafft FullEnrich ALLEIN die Kette Firma -> Entscheider ->
persoenliche Mail, Durchwahl, Mobilnummer fuer unsere echten Firmen?

Was hier NICHT passiert: kein Apollo, kein Hunter, kein Dropcontact,
kein Clay, kein LinkedIn. Keine Produktivdaten werden veraendert. Die
Ergebnisse liegen in einer eigenen Datenbank (daten/fullenrich-poc.db);
master.db, laeufe/ und kunden/ bleiben unberuehrt.

Reihenfolge - der Kern des Auftrags, weil sie ueber die Kosten entscheidet:

    Firmen aus unserer DB (deterministisch gezogen)
        |
    Webseite lesen: Startseite + Leistungs-/Loesungs-/Ueber-uns-Seiten
        |
    Automatisierungs-Klassifizierung (KI, strukturierte Antwort)
        |
    +-- YES       -> AUSGESCHLOSSEN, STOPP. Kein FullEnrich, kein Credit.
    +-- UNCERTAIN -> AUSGESCHLOSSEN, STOPP, zur Handpruefung.
    |
    +-- NO        -> FullEnrich company/lookup
                     -> people/search (CEO/GF/Inhaber/...)
                     -> Rangfolge + Firmen-Abgleich
                     -> contact/enrich/bulk
                     -> Pruefung
                     -> zulaessiger Kontakt
"""
import csv
import json
import random
import re
import sqlite3
import time
from datetime import datetime
from pathlib import Path

from pipeline.sources.fullenrich import (
    FullEnrichFehler, FullEnrichSource, KontingentLeer, beste_mail,
    ist_sammeladresse, telefon_art)

POC_DB = "daten/fullenrich-poc.db"
AUSWAHL_DATEI = "daten/fullenrich-poc-auswahl.json"
SEED = 20260824

# Titel fuer die Personensuche. Bewusst NICHT nur "CEO": bei deutschen
# Kleinfirmen heisst der Chef fast nie so.
TITEL_VARIANTEN = [
    "CEO", "Chief Executive Officer", "Geschäftsführer",
    "Geschäftsführerin", "Managing Director", "Owner", "Inhaber",
    "Inhaberin", "Geschäftsinhaber", "Founder", "Co-Founder", "Gründer",
    "Managing Partner", "Director", "Proprietor",
]

# Rangfolge laut Auftrag vom 24.08.2026. Weicht bewusst von
# pipeline/decision_maker.py ab (dort steht nach Olivers Vorgabe vom
# 19.08. der Inhaber vor dem CEO) - hier gilt die Vorgabe dieses Tests.
RANG_GRUPPEN = (
    ("ceo", "chief executive"),
    ("geschaftsfuhrer", "geschaeftsfuehrer", "geschaftsfuehrer",
     "geschaeftsfuhrer", "geschäftsführer"),
    ("owner", "inhaber", "geschaftsinhaber", "geschaeftsinhaber",
     "proprietor", "eigentumer", "eigentuemer"),
    ("managing director",),
    ("founder", "grunder", "gruender", "gründer"),
    ("managing partner",),
    ("director", "direktor", "vorstand", "geschaftsleitung",
     "geschaeftsleitung", "prokurist", "head of", "partner"),
)

SCHEMA = """
CREATE TABLE IF NOT EXISTS poc_ergebnisse (
    id INTEGER PRIMARY KEY,
    test_run_id TEXT NOT NULL,
    company_id INTEGER,
    company_kennung TEXT,
    company_name TEXT,
    domain TEXT,
    website TEXT,
    plz TEXT,
    ort TEXT,
    land TEXT,
    strasse TEXT,
    company_phone TEXT,
    company_email TEXT,
    sektor TEXT,
    keywords TEXT,
    mitarbeiter TEXT,
    ceo_owner_db TEXT,
    data_source TEXT,
    automation_status TEXT,
    automation_confidence REAL,
    automation_reason TEXT,
    automation_evidence TEXT,
    automation_source_url TEXT,
    automation_check_quelle TEXT,
    website_status TEXT,
    website_seiten TEXT,
    eligible INTEGER,
    exclusion_reason TEXT,
    fullenrich_company_found INTEGER,
    fullenrich_company_id TEXT,
    fullenrich_company_name TEXT,
    fullenrich_company_domain TEXT,
    company_match_status TEXT,
    company_match_score REAL,
    company_match_evidence TEXT,
    decision_maker_found INTEGER,
    decision_maker_first_name TEXT,
    decision_maker_last_name TEXT,
    decision_maker_name TEXT,
    decision_maker_title TEXT,
    decision_maker_headline TEXT,
    decision_maker_linkedin TEXT,
    decision_maker_rank INTEGER,
    decision_maker_reason TEXT,
    decision_maker_status TEXT,
    decision_maker_title_source TEXT,
    decision_maker_headline_evidence TEXT,
    alternative_candidates TEXT,
    individual_work_email TEXT,
    generic_company_email TEXT,
    private_personal_email TEXT,
    personal_email TEXT,
    personal_email_status TEXT,
    work_email TEXT,
    work_email_status TEXT,
    direct_phone TEXT,
    direct_phone_status TEXT,
    mobile_phone TEXT,
    mobile_phone_status TEXT,
    phones_raw TEXT,
    unknown_phones TEXT,
    work_email_is_generic INTEGER,
    verification_status TEXT,
    manual_review_reason TEXT,
    credits_used INTEGER,
    api_status TEXT,
    error_message TEXT,
    raw_response TEXT,
    created_at TEXT
);
CREATE TABLE IF NOT EXISTS poc_laeufe (
    test_run_id TEXT PRIMARY KEY,
    gestartet_am TEXT,
    beendet_am TEXT,
    seed INTEGER,
    firmen_geplant INTEGER,
    dry_run INTEGER,
    kennzahlen TEXT
);
"""


def log(*teile):
    print(f"[{datetime.now():%H:%M:%S}]", *teile, flush=True)


def _norm(text) -> str:
    wert = str(text or "").casefold().strip()
    for a, b in (("ä", "a"), ("ö", "o"), ("ü", "u"), ("ß", "ss")):
        wert = wert.replace(a, b)
    return wert


def _domain(wert) -> str:
    import re
    ohne = re.sub(r"^https?://", "", str(wert or "").strip().lower())
    return re.sub(r"^www\.", "", ohne).split("/")[0].strip()


def _ziffern(nummer) -> str:
    """Nationale Vergleichsform einer Rufnummer.

    "+49 5221 9660" und "05221 9660" muessen dieselbe Zeichenkette
    ergeben, sonst wird die Firmenzentrale nicht als solche erkannt.
    Entscheidend ist die SCHREIBWEISE des Originals: "+"/"00" heisst
    international (dann Laendervorwahl 49 abschneiden), sonst national
    (dann die fuehrende 0 abschneiden). An der reinen Ziffernfolge waere
    das nicht unterscheidbar - eine Ortsvorwahl wie 04941 faengt nach
    dem Nullstrich selbst mit 49 an."""
    import re
    roh = str(nummer or "").strip()
    kompakt = re.sub(r"\D", "", roh)
    if not kompakt:
        return ""
    if roh.startswith("+") or kompakt.startswith("00"):
        if kompakt.startswith("00"):
            kompakt = kompakt[2:]
        if kompakt.startswith("49"):
            kompakt = kompakt[2:]
        return kompakt.lstrip("0")
    return kompakt.lstrip("0")


# ------------------------------------------------------------- 1. Auswahl

def firmen_waehlen(daten_dir=".", limit=100, seed=SEED) -> list:
    """Deterministische Stichprobe aus der vorhandenen master.db.

    Es wird NICHTS eingefuegt und NICHTS geaendert - nur gelesen.
    Bevorzugt Firmen mit Name UND Domain: ohne Domain kann weder der
    Firmen-Abgleich noch die Personensuche sinnvoll laufen."""
    from pipeline.master_db import DB_NAME

    pfad = Path(daten_dir) / DB_NAME
    if not pfad.exists():
        raise FileNotFoundError(
            f"Master-Datenbank fehlt: {pfad}. Erst 'python -m pipeline "
            f"master-db' laufen lassen.")
    db = sqlite3.connect(pfad)
    db.row_factory = sqlite3.Row
    zeilen = db.execute(
        """SELECT id, kennung, name, domain, website, strasse, plz, ort,
                  land, telefon, email_allgemein, sektor, keywords,
                  mitarbeiter, ceo_owner
           FROM companies
           WHERE TRIM(COALESCE(name,'')) <> ''
             AND TRIM(COALESCE(domain,'')) <> ''
           ORDER BY kennung""").fetchall()
    db.close()

    gesamt = len(zeilen)
    if gesamt <= limit:
        gewaehlt = [dict(z) for z in zeilen]
    else:
        wuerfel = random.Random(seed)
        gewaehlt = [dict(z) for z in wuerfel.sample(list(zeilen), limit)]
        gewaehlt.sort(key=lambda f: f["kennung"])

    ziel = Path(daten_dir) / AUSWAHL_DATEI
    ziel.parent.mkdir(parents=True, exist_ok=True)
    ziel.write_text(json.dumps({
        "seed": seed, "limit": limit,
        "gueltige_firmen_in_db": gesamt,
        "gewaehlt": len(gewaehlt),
        "company_ids": [f["id"] for f in gewaehlt],
        "kennungen": [f["kennung"] for f in gewaehlt],
        "erstellt_am": datetime.now().isoformat(timespec="seconds"),
    }, ensure_ascii=False, indent=1), encoding="utf-8")
    return gewaehlt


# --------------------------------------------- 2. Webseite + Klassifizierung

def webseiten_und_urteile(firmen: list, ki, daten_dir=".", *,
                          ttl_tage=None, max_seiten=4, gleichzeitig=6,
                          seiten_leser=None) -> tuple:
    """(seiten_je_firma, urteile) - liest tief und urteilt strukturiert."""
    from concurrent.futures import ThreadPoolExecutor

    from pipeline import automation_llm
    from pipeline.website_tiefe import STANDARD_TTL_TAGE, seiten_lesen

    ttl_tage = STANDARD_TTL_TAGE if ttl_tage is None else ttl_tage
    leser = seiten_leser or (lambda website: seiten_lesen(
        website, daten_dir, max_seiten=max_seiten, ttl_tage=ttl_tage))

    log(f"    Webseiten lesen (Startseite + bis zu {max_seiten} Unterseiten, "
        f"Cache-TTL {ttl_tage} Tage) ...")
    seiten = {}

    def eine(paar):
        nummer, firma = paar
        return nummer, leser(firma.get("website") or firma.get("domain"))

    with ThreadPoolExecutor(
            max_workers=min(max(1, gleichzeitig), len(firmen) or 1)) as pool:
        for nummer, ergebnis in pool.map(eine, list(enumerate(firmen))):
            seiten[nummer] = ergebnis

    lesbar = sum(1 for e in seiten.values() if (e or {}).get("status") == "ok")
    log(f"    {lesbar} von {len(firmen)} Webseiten lieferten brauchbaren Text")

    log("    Automatisierungs-Klassifizierung ...")
    urteile = automation_llm.viele_klassifizieren(
        firmen, seiten, ki, gleichzeitig=gleichzeitig, log=log)
    return seiten, urteile


# ------------------------------------------------------ 3. Entscheider-Rang

def rang_von_titel(titel) -> int:
    """0 = bester Rang. len(RANG_GRUPPEN) = kein Entscheider-Titel."""
    text = _norm(titel)
    if text:
        for rang, begriffe in enumerate(RANG_GRUPPEN):
            if any(b in text for b in begriffe):
                return rang
    return len(RANG_GRUPPEN)


# --------------------------------------------- Titel aus der headline

# Nur echte Chef-Titel. Reihenfolge egal, aber die laengeren zuerst
# pruefen, damit "chief executive officer" nicht als "ceo" endet.
EXEC_TITEL = (
    "chief executive officer", "geschäftsführerin", "geschaeftsfuehrerin",
    "geschäftsführer", "geschaeftsfuehrer", "geschäftsinhaberin",
    "geschäftsinhaber", "managing director", "managing partner",
    "co-founder", "gründerin", "gruenderin", "gründer", "gruender",
    "inhaberin", "inhaber", "founder", "proprietor", "owner", "ceo",
)
# "<Titel> bei <Firma>" / "<Titel> at <Firma>". Ohne Trenner und ohne
# Firma zaehlt die headline NICHT - "Founder / entrepreneur" oder ein
# blosses "CEO" bleiben damit draussen.
_TRENNER = re.compile(r"\s+(?:bei|at|von|@|of)\s+", re.I)
# Rechtsformen und Fuellwoerter, die beim Firmenvergleich stoeren.
_RECHTSFORM = re.compile(
    r"(?i)\b(gmbh|mbh|ag|kg|ug|ohg|gbr|e\.?\s?k\.?|se|co|kgaa|"
    r"haftungsbeschr[aä]nkt|und|and|the)\b")


def _firmenkern(name) -> str:
    """Firmenname auf den Kern reduziert, damit
    "powerBridge Computervertrieb" und
    "powerBridge Computer Vertriebs GmbH" vergleichbar werden."""
    text = _norm(name)
    text = _RECHTSFORM.sub(" ", text)
    return re.sub(r"[^a-z0-9]", "", text)


def titel_aus_headline(headline, firmenname, firmen_domain="") -> dict:
    """{"titel", "firma_in_headline", "grund"} oder leeres titel.

    Regel vom 25.08.2026, bewusst eng: eine headline darf den Titel NUR
    dann belegen, wenn sie einen echten Chef-Titel traegt, der Form
    "<Titel> bei/at <Firma>" folgt UND die dort genannte Firma unsere
    ist. Alles andere - Selbstbeschreibungen, ein blosses "CEO", ein
    "Founder" ohne Firma - zaehlt weiterhin nicht.

    Grund fuer die Regel: FullEnrich fuellt employment[].title bei
    kleinen deutschen Firmen fast nie, schreibt die Funktion aber oft in
    die headline ("Geschäftsführer bei SYSCON GmbH"). Gemessen am
    25.08.2026: 3 von 5 zurueckgegebenen Personen."""
    leer = {"titel": "", "firma_in_headline": "", "grund": ""}
    text = (headline or "").strip()
    if not text:
        return leer

    niedrig = _norm(text)
    treffer = None
    for kandidat in EXEC_TITEL:
        stelle = niedrig.find(_norm(kandidat))
        if stelle >= 0:
            treffer = (kandidat, stelle)
            break
    if not treffer:
        return {**leer, "grund": "kein Chef-Titel in der headline"}

    teile = _TRENNER.split(text, maxsplit=1)
    if len(teile) < 2:
        return {**leer,
                "grund": f"Titel '{treffer[0]}' ohne 'bei/at <Firma>'"}
    kopf, firmenteil = teile[0], teile[1]
    # Der Titel muss VOR dem Trenner stehen, sonst ist er nicht die
    # Funktion dieser Person bei dieser Firma.
    if _norm(treffer[0]) not in _norm(kopf):
        return {**leer,
                "grund": f"Titel steht nicht vor 'bei/at' ({kopf[:40]})"}

    unser_kern = _firmenkern(firmenname)
    headline_kern = _firmenkern(firmenteil)
    if len(unser_kern) < 4:
        return {**leer, "grund": "unser Firmenname zu kurz zum Vergleich"}
    if unser_kern not in headline_kern and headline_kern not in unser_kern:
        return {**leer,
                "grund": f"headline nennt '{firmenteil[:45]}', "
                         f"wir suchen '{firmenname[:45]}'"}

    return {"titel": kopf.strip(" ,-|·"),
            "firma_in_headline": firmenteil.strip(),
            "grund": f"headline nennt '{kopf.strip()}' bei "
                     f"'{firmenteil.strip()[:45]}' - Firma stimmt überein"}


def _aktueller_titel(person: dict) -> tuple:
    """(titel, headline, firmenname, firmen_domain, linkedin).

    titel kommt AUSSCHLIESSLICH aus employment[].title. Die "headline"
    ist Selbstbeschreibung ("Digital problem solver") und wurde frueher
    faelschlich als Funktionsbezeichnung uebernommen - dadurch bekamen
    Leute ohne Entscheider-Rolle einen Titel angedichtet (Fund im echten
    Lauf 24.08.2026). Sie wird getrennt mitgefuehrt, nur zur Ansicht."""
    stellen = person.get("employment") or []
    aktuell = [s for s in stellen if isinstance(s, dict) and s.get("is_current")]
    stelle = (aktuell or [s for s in stellen if isinstance(s, dict)] or [{}])[0]
    firma = stelle.get("company")
    firmenname = firmen_domain = ""
    if isinstance(firma, dict):
        firmenname = firma.get("name") or ""
        firmen_domain = _domain(firma.get("domain") or firma.get("website"))
    elif isinstance(firma, str):
        firmenname = firma
    linkedin = ""
    profile = person.get("social_profiles")
    if isinstance(profile, dict):
        netz = profile.get("professional_network")
        if isinstance(netz, dict):
            linkedin = netz.get("url") or ""
    return ((stelle.get("title") or "").strip(),
            (person.get("headline") or "").strip(),
            firmenname, firmen_domain, linkedin)


def entscheider_waehlen(personen: list, firma: dict | None = None) -> tuple:
    """(gewaehlt, alternativen) - nie blind der erste Treffer.

    `firma` wird nur fuer die eng gefasste headline-Regel gebraucht
    (Auftrag 25.08.2026): fehlt employment[].title, darf der Titel aus
    der headline kommen - aber NUR in der Form "<Chef-Titel> bei
    <unsere Firma>". Ohne `firma` bleibt es beim alten Verhalten."""
    bewertet = []
    for person in personen or []:
        if not isinstance(person, dict):
            continue
        titel, headline, firmenname, firmen_domain, linkedin = \
            _aktueller_titel(person)
        titel_quelle = "employment.title" if titel else ""
        headline_beleg = ""
        if not titel and firma:
            aus_headline = titel_aus_headline(
                headline, firma.get("name"), firmen_domain)
            if aus_headline["titel"]:
                titel = aus_headline["titel"]
                titel_quelle = "headline (Firma stimmt überein)"
                headline_beleg = aus_headline["grund"]
        bewertet.append({
            "titel_quelle": titel_quelle,
            "headline_beleg": headline_beleg,
            "name": person.get("full_name")
                    or f"{person.get('first_name','')} "
                       f"{person.get('last_name','')}".strip(),
            "first_name": person.get("first_name", ""),
            "last_name": person.get("last_name", ""),
            "titel": titel, "headline": headline,
            "firmenname": firmenname, "firmen_domain": firmen_domain,
            "linkedin": linkedin, "rang": rang_von_titel(titel),
        })
    if not bewertet:
        return None, []
    bewertet.sort(key=lambda p: p["rang"])
    bester = dict(bewertet[0])
    if bester["rang"] >= len(RANG_GRUPPEN):
        zusatz = (f" (headline: {bester['headline'][:60]})"
                  if bester["headline"] else "")
        bester["rang_grund"] = f"kein Entscheider-Titel erkannt{zusatz}"
    else:
        bester["rang_grund"] = f"Rang {bester['rang']} ({bester['titel']})"
    return bester, bewertet[1:6]


# --------------------------------------------------- 4. Firmen-Abgleich

def firmen_abgleich(unsere: dict, fe_firma: dict, person: dict) -> tuple:
    """(status, score, belege) - damit keine Person zur falschen Firma faellt."""
    belege, punkte = [], 0.0
    unsere_domain = _domain(unsere.get("domain") or unsere.get("website"))
    unser_name = _norm(unsere.get("name"))
    unser_ort = _norm(unsere.get("ort"))

    fe_domain = fe_name = fe_ort = ""
    if isinstance(fe_firma, dict):
        fe_domain = _domain(fe_firma.get("domain") or fe_firma.get("website"))
        fe_name = _norm(fe_firma.get("name"))
        adresse = fe_firma.get("location") or fe_firma.get("address") or {}
        if isinstance(adresse, dict):
            fe_ort = _norm(adresse.get("city"))

    person_domain = _norm(person.get("firmen_domain")) if person else ""
    person_firma = _norm(person.get("firmenname")) if person else ""

    if unsere_domain and fe_domain:
        if unsere_domain == fe_domain:
            punkte += 0.5
            belege.append(f"Domain identisch ({unsere_domain})")
        else:
            belege.append(f"Domain weicht ab (wir {unsere_domain}, "
                          f"FullEnrich {fe_domain})")
            return "no_match", 0.0, "; ".join(belege)

    if unsere_domain and person_domain:
        if unsere_domain == person_domain:
            punkte += 0.3
            belege.append("Person sitzt laut Profil auf derselben Domain")
        else:
            belege.append(f"Person-Domain weicht ab ({person_domain})")

    if unser_name and fe_name:
        if unser_name == fe_name:
            punkte += 0.2
            belege.append("Firmenname identisch")
        elif unser_name in fe_name or fe_name in unser_name:
            punkte += 0.1
            belege.append("Firmenname teilweise gleich")

    if unser_name and person_firma:
        if unser_name == person_firma or unser_name in person_firma \
                or person_firma in unser_name:
            punkte += 0.2
            belege.append("Person nennt dieselbe Firma")
        else:
            belege.append(f"Person nennt andere Firma ({person_firma[:40]})")

    if unser_ort and fe_ort and unser_ort == fe_ort:
        punkte += 0.1
        belege.append(f"Ort stimmt ({fe_ort})")

    punkte = round(min(punkte, 1.0), 2)
    if punkte >= 0.7:
        status = "strong_match"
    elif punkte >= 0.4:
        status = "possible_match"
    elif punkte > 0:
        status = "uncertain"
    else:
        status = "uncertain" if (fe_firma or person) else "no_match"
    return status, punkte, "; ".join(belege) or "keine vergleichbaren Felder"


def pruefung_kontakt(zeile: dict) -> tuple:
    """(verification_status, grund_fuer_handpruefung).

    strong_match  Firma sicher, Entscheider-Titel echt, Mail brauchbar
    possible_match eines davon schwaecher
    uncertain      mehrere Zweifel - nicht stillschweigend als geprueft
                   speichern
    invalid        gar keine brauchbare Kontaktangabe
    """
    gruende = []
    if not zeile["decision_maker_found"]:
        return "invalid", "kein Entscheider gefunden"
    if not (zeile["personal_email"] or zeile["work_email"]
            or zeile["mobile_phone"] or zeile["direct_phone"]):
        return "invalid", "keine Kontaktangabe gefunden"

    if zeile["company_match_status"] == "no_match":
        return "invalid", "Person gehoert nachweislich zu anderer Firma"
    if zeile["company_match_status"] == "uncertain":
        gruende.append("Firmen-Zuordnung unsicher")
    if (zeile["decision_maker_rank"] is None
            or zeile["decision_maker_rank"] >= len(RANG_GRUPPEN)):
        gruende.append("kein echter Entscheider-Titel")
    if zeile["work_email_status"] == "CATCH_ALL" and not zeile["personal_email"]:
        gruende.append("nur Catch-all-Mail (Server nimmt alles an)")

    if not gruende and zeile["company_match_status"] == "strong_match":
        return "strong_match", ""
    if len(gruende) <= 1:
        return "possible_match", "; ".join(gruende)
    return "uncertain", "; ".join(gruende)


# --------------------------------------------------------- 5. Speicherung

def db_oeffnen(daten_dir="."):
    pfad = Path(daten_dir) / POC_DB
    pfad.parent.mkdir(parents=True, exist_ok=True)
    db = sqlite3.connect(pfad)
    db.executescript(SCHEMA)
    return db


def _saubern(roh) -> str:
    """Rohantwort fuer die Fehlersuche - ohne Schluessel, ohne Header."""
    if roh is None:
        return ""
    try:
        text = json.dumps(roh, ensure_ascii=False)
    except (TypeError, ValueError):
        text = str(roh)
    for wort in ("authorization", "bearer", "api_key", "apikey", "token"):
        if wort in text.lower():
            return json.dumps({"redigiert": "Antwort enthielt ein Feld, das "
                                            "wie ein Zugangsschluessel aussieht"})
    return text[:20000]


def ergebnis_speichern(db, satz: dict):
    spalten = [s for s in satz if s != "id"]
    platz = ",".join("?" * len(spalten))
    db.execute(f"INSERT INTO poc_ergebnisse ({','.join(spalten)}) "
               f"VALUES ({platz})", [satz[s] for s in spalten])
    db.commit()


# ------------------------------------------------------------- 6. Kennzahlen

def kennzahlen_rechnen(zeilen: list) -> dict:
    def anteil(zaehler, nenner):
        return round(100.0 * zaehler / nenner, 1) if nenner else 0.0

    getestet = len(zeilen)
    ja = [z for z in zeilen if z["automation_status"] == "YES"]
    nein = [z for z in zeilen if z["automation_status"] == "NO"]
    unklar = [z for z in zeilen if z["automation_status"] == "UNCERTAIN"]
    zulaessig = [z for z in zeilen if z["eligible"]]

    gematcht = [z for z in zulaessig if z["fullenrich_company_found"]]
    stark = [z for z in zulaessig
             if z["company_match_status"] == "strong_match"]
    leute = [z for z in zulaessig if z["decision_maker_found"]]
    entscheider = [z for z in leute
                   if z.get("decision_maker_status") == "valid"]
    unsichere_leute = [z for z in leute
                       if z.get("decision_maker_status") != "valid"]
    echter_titel = [z for z in leute
                    if z["decision_maker_rank"] is not None
                    and z["decision_maker_rank"] < len(RANG_GRUPPEN)]
    individuell = [z for z in zulaessig if z.get("individual_work_email")]
    generisch = [z for z in zulaessig if z.get("generic_company_email")]
    privat = [z for z in zulaessig if z.get("private_personal_email")]
    persoenlich = [z for z in zulaessig if z["personal_email"]]
    arbeit = [z for z in zulaessig if z["work_email"]]
    beide = [z for z in zulaessig if z["personal_email"] and z["work_email"]]
    keine_mail = [z for z in zulaessig
                  if not z["personal_email"] and not z["work_email"]]
    durchwahl = [z for z in zulaessig if z["direct_phone"]]
    mobil = [z for z in zulaessig if z["mobile_phone"]]
    # "Vollstaendig" nach Olivers Anforderung: richtige Firma, echter
    # Entscheider, mindestens eine Mail UND mindestens eine persoenliche
    # Nummer. Das ist streng - genau darum geht es.
    # Brauchbar heisst: glaubwuerdiger Entscheider UND eine
    # personenbezogene Geschaefts-Mail. info@ zaehlt hier NICHT.
    brauchbar = [z for z in entscheider if z.get("individual_work_email")]
    vollstaendig = [z for z in brauchbar
                    if z["mobile_phone"] or z["direct_phone"]]
    handpruefung = [z for z in zeilen
                    if z["manual_review_reason"]
                    or z["automation_status"] == "UNCERTAIN"
                    or (z["eligible"] and z["verification_status"] == "uncertain")]
    fehler = [z for z in zeilen if str(z["api_status"]).startswith("error")]

    credits = sum(int(z["credits_used"] or 0) for z in zeilen)
    return {
        "companies_tested": getestet,
        "automation_yes": len(ja),
        "automation_no": len(nein),
        "automation_uncertain": len(unklar),
        "eligible_companies": len(zulaessig),
        "companies_skipped_before_enrichment": len(ja) + len(unklar),
        "automation_exclusion_rate": anteil(len(ja), getestet),
        "companies_matched": len(gematcht),
        "company_match_rate": anteil(len(gematcht), len(zulaessig)),
        "strong_company_matches": len(stark),
        "people_found": len(leute),
        "credible_decision_makers": len(entscheider),
        "uncertain_people": len(unsichere_leute),
        "decision_makers_found": len(entscheider),
        "decision_maker_rate": anteil(len(entscheider), len(zulaessig)),
        "real_decision_maker_titles": len(echter_titel),
        "real_decision_maker_rate": anteil(len(echter_titel), len(zulaessig)),
        "individual_work_emails": len(individuell),
        "individual_work_email_rate": anteil(len(individuell), len(zulaessig)),
        "generic_company_emails": len(generisch),
        "private_personal_emails": len(privat),
        "personal_emails_found": len(persoenlich),
        "personal_email_rate": anteil(len(persoenlich), len(zulaessig)),
        "work_emails_found": len(arbeit),
        "work_email_rate": anteil(len(arbeit), len(zulaessig)),
        "both_emails_found": len(beide),
        "no_email_found": len(keine_mail),
        "direct_phones_found": len(durchwahl),
        "direct_phone_rate": anteil(len(durchwahl), len(zulaessig)),
        "mobile_numbers_found": len(mobil),
        "mobile_rate": anteil(len(mobil), len(zulaessig)),
        "usable_contacts": len(brauchbar),
        "usable_contact_rate": anteil(len(brauchbar), len(zulaessig)),
        "fully_enriched_contacts": len(vollstaendig),
        "fully_enriched_rate": anteil(len(vollstaendig), len(zulaessig)),
        "manual_review_count": len(handpruefung),
        "api_failures": len(fehler),
        "total_credits_used": credits,
        "average_credits_per_company": round(credits / getestet, 2)
        if getestet else 0.0,
        "average_credits_per_eligible_company":
            round(credits / len(zulaessig), 2) if zulaessig else 0.0,
        "average_credits_per_usable_contact":
            round(credits / len(brauchbar), 2) if brauchbar else 0.0,
        "match_status_verteilung": {
            status: sum(1 for z in zulaessig
                        if z["company_match_status"] == status)
            for status in ("strong_match", "possible_match", "uncertain",
                           "no_match")},
    }


def bericht_text(k: dict) -> str:
    z = ["=" * 52, "FULLENRICH-ONLY BENCHMARK", "=" * 52, "",
         f"Companies tested:                  {k['companies_tested']:>6}", "",
         "-- Automation gate (before any credit) --",
         f"Automation YES (excluded):         {k['automation_yes']:>6}"
         f"  {k['automation_exclusion_rate']:>5}%",
         f"Automation NO  (eligible):         {k['automation_no']:>6}",
         f"Automation UNCERTAIN (review):     {k['automation_uncertain']:>6}",
         f"Skipped before enrichment:         "
         f"{k['companies_skipped_before_enrichment']:>6}",
         f"Eligible companies:                {k['eligible_companies']:>6}",
         "", "-- FullEnrich (eligible only) --",
         f"Companies matched:                 {k['companies_matched']:>6}"
         f"  {k['company_match_rate']:>5}%",
         f"  of these strong matches:         {k['strong_company_matches']:>6}",
         f"People returned by search:         {k.get('people_found', 0):>6}",
         f"CREDIBLE decision makers:          "
         f"{k.get('credible_decision_makers', 0):>6}"
         f"  {k['decision_maker_rate']:>5}%",
         f"  uncertain people (not counted):   "
         f"{k.get('uncertain_people', 0):>6}",
         f"  with a real decision-maker title:{k['real_decision_maker_titles']:>6}"
         f"  {k['real_decision_maker_rate']:>5}%",
         "", "-- E-Mail (three separate categories) --",
         f"INDIVIDUAL work emails:            "
         f"{k.get('individual_work_emails', 0):>6}"
         f"  {k.get('individual_work_email_rate', 0):>5}%   <- Oliver's target",
         f"Generic company emails (info@ ...):"
         f"{k.get('generic_company_emails', 0):>6}   not counted as a contact",
         f"Private personal emails (gmail...):"
         f"{k.get('private_personal_emails', 0):>6}",
         f"Work emails total:                 {k['work_emails_found']:>6}"
         f"  {k['work_email_rate']:>5}%",
         f"Neither found:                     {k['no_email_found']:>6}",
         "", "-- Phone (company switchboard never counts) --",
         f"Direct lines found:                {k['direct_phones_found']:>6}"
         f"  {k['direct_phone_rate']:>5}%",
         f"Mobile numbers found:              {k['mobile_numbers_found']:>6}"
         f"  {k['mobile_rate']:>5}%",
         "", "-- Result --",
         f"Usable contacts (name+mail):       {k['usable_contacts']:>6}"
         f"  {k['usable_contact_rate']:>5}%",
         f"Fully enriched (mail AND phone):   "
         f"{k['fully_enriched_contacts']:>6}  {k['fully_enriched_rate']:>5}%",
         f"Needing manual review:             {k['manual_review_count']:>6}",
         f"FullEnrich API failures:           {k['api_failures']:>6}",
         "", "-- Company match quality --"]
    for status, anzahl in k["match_status_verteilung"].items():
        z.append(f"  {status:<24} {anzahl:>6}")
    z += ["", "-- Credits --",
          f"Total credits used:                {k['total_credits_used']:>6}",
          f"Average per company:               "
          f"{k['average_credits_per_company']:>6}",
          f"Average per eligible company:      "
          f"{k['average_credits_per_eligible_company']:>6}",
          f"Average per usable contact:        "
          f"{k['average_credits_per_usable_contact']:>6}"]
    if "estimated_search_credits" in k:
        z.append(f"Search credits (estimated):        "
                 f"{k['estimated_search_credits']:>6}"
                 f"   0.25/hit, not reported by the API")
    if k.get("estimated_credits_min") is not None:
        z += ["",
              f"ESTIMATED credits for a real run:",
              f"  minimum (searches only):         "
              f"{k['estimated_credits_min']:>6}",
              f"  maximum (all with mail+mobile):  "
              f"{k['estimated_credits_max']:>6}"]
    if k.get("credits_measured") is not None:
        z += ["",
              f"Credits before run:                {k['credits_before']:>6}",
              f"Credits after run:                 {k['credits_after']:>6}",
              f"MEASURED credits used:             {k['credits_measured']:>6}",
              f"  per eligible company:            "
              f"{k['measured_credits_per_eligible_company']:>6}",
              f"  per decision maker:              "
              f"{k['measured_credits_per_decision_maker']:>6}",
              f"  per usable contact:              "
              f"{k['measured_credits_per_usable_contact']:>6}"]
    z.append("=" * 52)
    return "\n".join(z)


def credits_schaetzen(zulaessige: int) -> dict:
    """Was ein echter Lauf ungefaehr kosten wuerde - als SPANNE.

    Genau ist das nicht moeglich: die Suche kostet je Treffer, und wie
    viele Treffer kommen, weiss man erst nach dem (bezahlten) Aufruf.
    Deshalb zwei Ecken, beide aus der Preisliste der Doku gerechnet:

      Untergrenze  nur die Suchen, niemand wird angereichert
                   (Firmentreffer 0,25 + ~2 Personen a 0,25)
      Obergrenze   jede zulaessige Firma liefert einen glaubwuerdigen
                   Entscheider MIT Mail und Mobilnummer
                   (Suchen + 1 Mail + 10 Mobil)

    Die Wahrheit liegt dazwischen und haengt an der Trefferquote, die
    dieser Test ja gerade erst messen soll."""
    from pipeline.sources.fullenrich import (
        PREIS_MAIL, PREIS_MOBIL, PREIS_SUCHTREFFER)

    such_anteil = PREIS_SUCHTREFFER * 3          # 1 Firma + ~2 Personen
    unten = round(zulaessige * such_anteil, 2)
    oben = round(zulaessige * (such_anteil + PREIS_MAIL + PREIS_MOBIL), 2)
    return {"estimated_credits_min": unten,
            "estimated_credits_max": oben,
            "estimated_credits_note":
                "Untergrenze = nur Suchen; Obergrenze = jede Firma mit "
                "Mail und Mobilnummer. Preise laut FullEnrich-Doku."}


# ---------------------------------------------------------------- 7. Export

EXPORT_SPALTEN = [
    "company_id", "data_source", "company", "domain", "website",
    "strasse", "plz", "ort", "land",
    "sektor", "keywords", "mitarbeiter", "ceo_owner_db",
    "company_phone", "company_email",
    "automation_status", "automation_confidence", "automation_reason",
    "automation_evidence", "automation_source_url", "website_status",
    "eligible", "exclusion_reason",
    "fullenrich_company_found", "fullenrich_company_name",
    "company_match_status", "company_match_score",
    "decision_maker", "title", "headline", "linkedin",
    "decision_maker_status", "title_source", "headline_evidence",
    "individual_work_email", "generic_company_email",
    "private_personal_email",
    "personal_email", "personal_email_status",
    "work_email", "work_email_status", "work_email_is_generic",
    "direct_line", "direct_line_status",
    "mobile", "mobile_status", "unknown_phones",
    "verification_status", "manual_review_reason",
    "credits_used", "api_status", "error",
]


def _export_zeile(z: dict) -> list:
    return [
        z["company_id"], z.get("data_source", ""), z["company_name"],
        z["domain"], z["website"], z.get("strasse", ""),
        z["plz"], z["ort"], z["land"],
        z.get("sektor", ""), z.get("keywords", ""), z.get("mitarbeiter", ""),
        z.get("ceo_owner_db", ""),
        z["company_phone"], z.get("company_email", ""),
        z["automation_status"], z["automation_confidence"],
        (z["automation_reason"] or "")[:300],
        (z["automation_evidence"] or "")[:300],
        z["automation_source_url"], z["website_status"],
        "yes" if z["eligible"] else "no", z["exclusion_reason"],
        "yes" if z["fullenrich_company_found"] else "no",
        z["fullenrich_company_name"],
        z["company_match_status"], z["company_match_score"],
        z["decision_maker_name"], z["decision_maker_title"],
        (z["decision_maker_headline"] or "")[:120], z["decision_maker_linkedin"],
        z.get("decision_maker_status", ""),
        z.get("decision_maker_title_source", ""),
        (z.get("decision_maker_headline_evidence") or "")[:150],
        z.get("individual_work_email", ""),
        z.get("generic_company_email", ""),
        z.get("private_personal_email", ""),
        z["personal_email"], z["personal_email_status"],
        z["work_email"], z["work_email_status"],
        "yes" if z.get("work_email_is_generic") else "no",
        z["direct_phone"], z["direct_phone_status"],
        z["mobile_phone"], z["mobile_phone_status"],
        z.get("unknown_phones", ""),
        z["verification_status"], z["manual_review_reason"],
        z["credits_used"], z["api_status"], z["error_message"],
    ]


def _csv_schreiben(pfad: Path, zeilen: list):
    with pfad.open("w", newline="", encoding="utf-8") as f:
        schreiber = csv.writer(f)
        schreiber.writerow(EXPORT_SPALTEN)
        for zeile in zeilen:
            schreiber.writerow(_export_zeile(zeile))


def handpruefung_liste(zeilen: list) -> list:
    """Alles, was ein Mensch anschauen muss - nichts wird still verworfen."""
    treffer = []
    for zeile in zeilen:
        gruende = []
        if zeile["automation_status"] == "UNCERTAIN":
            gruende.append("Automatisierung unklar")
        if zeile["eligible"]:
            if zeile["company_match_status"] in ("uncertain", "no_match"):
                gruende.append("Firmen-Zuordnung unsicher")
            if zeile["verification_status"] == "uncertain":
                gruende.append("Kontakt unsicher")
            if zeile["manual_review_reason"]:
                gruende.append(zeile["manual_review_reason"])
        if gruende:
            treffer.append((" | ".join(dict.fromkeys(gruende)), zeile))
    return treffer


def exportieren(zeilen: list, daten_dir=".", stempel=None) -> dict:
    """Drei Dateien: alles, nur zulaessige Kontakte, Handpruefung."""
    stempel = stempel or datetime.now().strftime("%Y%m%d-%H%M")
    basis = Path(daten_dir)
    ziele = {}

    alle = basis / f"fullenrich-alle-{stempel}.csv"
    _csv_schreiben(alle, zeilen)
    ziele["alle_csv"] = alle

    kontakte = [z for z in zeilen
                if z["eligible"]
                and z.get("decision_maker_status") == "valid"
                and z.get("individual_work_email")]
    kontakt_pfad = basis / f"fullenrich-kontakte-{stempel}.csv"
    _csv_schreiben(kontakt_pfad, kontakte)
    ziele["kontakte_csv"] = kontakt_pfad

    pruefung = handpruefung_liste(zeilen)
    pruef_pfad = basis / f"fullenrich-handpruefung-{stempel}.csv"
    with pruef_pfad.open("w", newline="", encoding="utf-8") as f:
        schreiber = csv.writer(f)
        schreiber.writerow(["grund"] + EXPORT_SPALTEN)
        for grund, zeile in pruefung:
            schreiber.writerow([grund] + _export_zeile(zeile))
    ziele["handpruefung_csv"] = pruef_pfad

    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill

        wb = openpyxl.Workbook()
        blaetter = (("Alle Firmen", [(None, z) for z in zeilen]),
                    ("Zulässige Kontakte", [(None, z) for z in kontakte]),
                    ("Handprüfung", pruefung))
        for nummer, (titel, daten) in enumerate(blaetter):
            blatt = wb.active if nummer == 0 else wb.create_sheet()
            blatt.title = titel
            mit_grund = titel == "Handprüfung"
            blatt.append((["Grund"] if mit_grund else []) + EXPORT_SPALTEN)
            for grund, zeile in daten:
                blatt.append(([grund] if mit_grund else [])
                             + _export_zeile(zeile))
            for zelle in blatt[1]:
                zelle.font = Font(bold=True, color="FFFFFFFF", size=10)
                zelle.fill = PatternFill("solid", fgColor="FF1F3A56")
                zelle.alignment = Alignment(vertical="center", wrap_text=True)
            blatt.freeze_panes = "A2"
            if blatt.max_row > 1:
                blatt.auto_filter.ref = blatt.dimensions
        xlsx = basis / f"fullenrich-{stempel}.xlsx"
        wb.save(xlsx)
        ziele["xlsx"] = xlsx
    except ImportError:
        pass
    return ziele


# ------------------------------------------------------------- 8. Der Lauf

def _leere_zeile(firma, test_run_id) -> dict:
    return {
        "test_run_id": test_run_id,
        "company_id": firma.get("id"),
        "company_kennung": firma.get("kennung", ""),
        "company_name": firma.get("name", ""),
        "domain": _domain(firma.get("domain") or firma.get("website")),
        "website": firma.get("website", ""),
        "plz": firma.get("plz", ""), "ort": firma.get("ort", ""),
        "land": firma.get("land", ""),
        "strasse": firma.get("strasse", ""),
        "company_phone": firma.get("telefon", ""),
        "company_email": firma.get("email_allgemein", ""),
        "sektor": firma.get("sektor", ""),
        "keywords": firma.get("keywords", ""),
        "mitarbeiter": firma.get("mitarbeiter", ""),
        "ceo_owner_db": firma.get("ceo_owner", ""),
        # Woher die Firma ueberhaupt stammt - fuer die Frage nach der
        # Markt-Abdeckung (Punkt 6) getrennt von der Anreicherung.
        "data_source": firma.get("data_source", "master.db"),
        "automation_status": "", "automation_confidence": 0.0,
        "automation_reason": "", "automation_evidence": "",
        "automation_source_url": "", "automation_check_quelle": "",
        "website_status": "", "website_seiten": "",
        "eligible": 0, "exclusion_reason": "",
        "fullenrich_company_found": 0, "fullenrich_company_id": "",
        "fullenrich_company_name": "", "fullenrich_company_domain": "",
        "company_match_status": "", "company_match_score": 0.0,
        "company_match_evidence": "",
        "decision_maker_found": 0, "decision_maker_first_name": "",
        "decision_maker_last_name": "", "decision_maker_name": "",
        "decision_maker_title": "", "decision_maker_headline": "",
        "decision_maker_linkedin": "", "decision_maker_rank": None,
        "decision_maker_reason": "", "decision_maker_status": "",
        "decision_maker_title_source": "",
        "decision_maker_headline_evidence": "",
        "alternative_candidates": "",
        "individual_work_email": "", "generic_company_email": "",
        "private_personal_email": "",
        "personal_email": "", "personal_email_status": "",
        "work_email": "", "work_email_status": "",
        "direct_phone": "", "direct_phone_status": "",
        "mobile_phone": "", "mobile_phone_status": "",
        "phones_raw": "", "unknown_phones": "",
        "work_email_is_generic": 0, "verification_status": "",
        "manual_review_reason": "",
        "credits_used": 0, "api_status": "", "error_message": "",
        "raw_response": "",
        "created_at": datetime.now().isoformat(timespec="seconds"),
    }


def lauf(daten_dir=".", limit=100, api_key=None, quelle=None, ki=None,
         seed=SEED, roh_speichern=True, max_credits=None, dry_run=False,
         ttl_tage=None, seiten_leser=None, fe_gleichzeitig=1) -> dict:
    """Der ganze Lauf. dry_run=True macht KEINEN FullEnrich-Aufruf."""
    import os

    test_run_id = f"fe-{datetime.now():%Y%m%d-%H%M%S}"
    start = time.time()

    firmen = firmen_waehlen(daten_dir, limit, seed)
    print("FullEnrich-only enrichment")
    print(f"Companies selected for test: {len(firmen)}")
    if dry_run:
        print("DRY RUN - kein einziger FullEnrich-Aufruf, keine Credits")
    log(f"Auswahl deterministisch (seed {seed}), gespeichert in "
        f"{AUSWAHL_DATEI}")

    if ki is None and not dry_run or ki is None:
        try:
            from pipeline.ki import KI
            ki = KI()
        except Exception as fehler:      # noqa: BLE001
            log(f"KI nicht verfügbar ({fehler}) - alle Firmen gelten als "
                f"UNCERTAIN und werden nicht angereichert.")
            ki = None

    log("Schritt 1/3: Webseiten lesen und Automatisierung klassifizieren "
        "(VOR jedem FullEnrich-Credit)")
    seiten, urteile = webseiten_und_urteile(
        firmen, ki, daten_dir, ttl_tage=ttl_tage, seiten_leser=seiten_leser)

    guthaben_vorher = None
    if not dry_run:
        quelle = quelle or FullEnrichSource(
            api_key or os.environ.get("FULLENRICH_API_KEY"))
        try:
            guthaben_vorher = quelle.guthaben()
            log(f"    Guthaben vor dem Lauf: {guthaben_vorher}")
        except FullEnrichFehler as fehler:
            log(f"    Guthaben nicht lesbar: {fehler}")

    db = db_oeffnen(daten_dir)
    db.execute("INSERT OR REPLACE INTO poc_laeufe (test_run_id, gestartet_am,"
               " seed, firmen_geplant, dry_run) VALUES (?,?,?,?,?)",
               (test_run_id, datetime.now().isoformat(timespec="seconds"),
                seed, len(firmen), 1 if dry_run else 0))
    db.commit()

    zeilen = []
    kontingent_leer = False
    such_credits = [0.0]
    log("Schritt 2/3: FullEnrich für die zulässigen Firmen")
    for nummer, firma in enumerate(firmen):
        zeile = _leere_zeile(firma, test_run_id)
        urteil = urteile.get(nummer) or {}
        seiten_info = seiten.get(nummer) or {}
        zeile["automation_status"] = urteil.get("automation_status",
                                                "UNCERTAIN")
        zeile["automation_confidence"] = urteil.get("confidence", 0.0)
        zeile["automation_reason"] = urteil.get("reason", "")
        zeile["automation_evidence"] = urteil.get("evidence", "")
        zeile["automation_source_url"] = urteil.get("source_url", "")
        zeile["automation_check_quelle"] = urteil.get("quelle", "")
        zeile["website_status"] = seiten_info.get("status", "")
        zeile["website_seiten"] = json.dumps(
            [s["url"] for s in (seiten_info.get("seiten") or [])],
            ensure_ascii=False)

        # ---- HARTE REGEL: hier endet es fuer YES und UNCERTAIN.
        if zeile["automation_status"] == "YES":
            zeile["exclusion_reason"] = "automation_provider"
            zeile["api_status"] = "skipped_no_api_call"
            zeilen.append(zeile)
            ergebnis_speichern(db, zeile)
            continue
        if zeile["automation_status"] != "NO":
            zeile["exclusion_reason"] = "automation_uncertain"
            zeile["api_status"] = "skipped_no_api_call"
            zeile["manual_review_reason"] = urteil.get("reason", "")[:200]
            zeilen.append(zeile)
            ergebnis_speichern(db, zeile)
            continue

        zeile["eligible"] = 1
        if dry_run:
            zeile["api_status"] = "dry_run"
            zeilen.append(zeile)
            ergebnis_speichern(db, zeile)
            continue
        if kontingent_leer:
            zeile["api_status"] = "aborted_credits_insufficient"
            zeile["error_message"] = "Lauf nach Guthaben-Ende abgebrochen"
            zeilen.append(zeile)
            ergebnis_speichern(db, zeile)
            continue

        verbraucht = (sum(int(z["credits_used"] or 0) for z in zeilen)
                      + such_credits[0])
        if max_credits is not None and verbraucht >= max_credits:
            zeile["api_status"] = "budget_exhausted"
            zeile["error_message"] = (
                f"Credit-Bremse erreicht ({verbraucht:.2f} von {max_credits})")
            zeilen.append(zeile)
            ergebnis_speichern(db, zeile)
            continue

        try:
            _eine_firma(quelle, firma, zeile, roh_speichern, such_credits)
        except KontingentLeer as fehler:
            kontingent_leer = True
            zeile["api_status"] = "credits_insufficient"
            zeile["error_message"] = str(fehler)[:400]
            log(f"ABBRUCH: {fehler}")
        except FullEnrichFehler as fehler:
            zeile["api_status"] = f"error_{fehler.status or 'unknown'}"
            zeile["error_message"] = str(fehler)[:400]
        except Exception as fehler:      # noqa: BLE001 - eine Firma darf
            # den ganzen Lauf nicht kippen
            zeile["api_status"] = "error_unexpected"
            zeile["error_message"] = f"{type(fehler).__name__}: {fehler}"[:400]

        zustand, grund = pruefung_kontakt(zeile)
        zeile["verification_status"] = zustand
        if grund and not zeile["manual_review_reason"]:
            zeile["manual_review_reason"] = grund

        zeilen.append(zeile)
        ergebnis_speichern(db, zeile)
        if (nummer + 1) % 10 == 0:
            log(f"    {nummer + 1}/{len(firmen)} Firmen verarbeitet")

    log("Schritt 3/3: Auswertung und Export")
    guthaben_nachher = None
    if not dry_run and quelle is not None:
        try:
            guthaben_nachher = quelle.guthaben()
        except FullEnrichFehler as fehler:
            log(f"    Guthaben nicht lesbar: {fehler}")
    kennzahlen = kennzahlen_rechnen(zeilen)
    kennzahlen["estimated_search_credits"] = round(such_credits[0], 2)
    kennzahlen["credits_before"] = guthaben_vorher
    kennzahlen["credits_after"] = guthaben_nachher
    if dry_run:
        kennzahlen.update(credits_schaetzen(kennzahlen["eligible_companies"]))
    if guthaben_vorher is not None and guthaben_nachher is not None:
        # Der GEMESSENE Verbrauch - inklusive der Such-Credits, die die
        # Anreicherungs-Antwort nicht mitmeldet. Das ist die ehrliche
        # Zahl; alles andere sind Schaetzungen.
        verbrauch = round(guthaben_vorher - guthaben_nachher, 2)
        kennzahlen["credits_measured"] = verbrauch
        zul = kennzahlen["eligible_companies"] or 0
        dm = kennzahlen["decision_makers_found"] or 0
        nutz = kennzahlen["usable_contacts"] or 0
        kennzahlen["measured_credits_per_eligible_company"] = (
            round(verbrauch / zul, 2) if zul else 0.0)
        kennzahlen["measured_credits_per_decision_maker"] = (
            round(verbrauch / dm, 2) if dm else 0.0)
        kennzahlen["measured_credits_per_usable_contact"] = (
            round(verbrauch / nutz, 2) if nutz else 0.0)
    db.execute("UPDATE poc_laeufe SET beendet_am=?, kennzahlen=? "
               "WHERE test_run_id=?",
               (datetime.now().isoformat(timespec="seconds"),
                json.dumps(kennzahlen, ensure_ascii=False), test_run_id))
    db.commit()
    db.close()

    ziele = exportieren(zeilen, daten_dir)
    return {"test_run_id": test_run_id, "kennzahlen": kennzahlen,
            "zeilen": zeilen, "exporte": ziele, "dry_run": dry_run,
            "laufzeit_minuten": round((time.time() - start) / 60, 1)}


def _entscheider_guete(zeile: dict, person: dict) -> str:
    """"valid" | "uncertain" | "company_mismatch".

    Nur "valid" oeffnet das Tor zur bezahlten Anreicherung. Verlangt
    wird BEIDES: die Firma muss sicher zugeordnet sein UND die Person
    muss einen echten Entscheider-Titel tragen. Eine headline
    ("Digital problem solver") ist kein Titel - genau dafuer wurde
    dieses Tor gebaut."""
    if zeile["company_match_status"] == "no_match":
        return "company_mismatch"
    if zeile["company_match_status"] != "strong_match":
        return "uncertain"
    if not (person.get("titel") or "").strip():
        return "uncertain"
    if person.get("rang") is None or person["rang"] >= len(RANG_GRUPPEN):
        return "uncertain"
    # Die Person muss auch laut ihrem eigenen Profil zu dieser Firma
    # gehoeren, sonst ist die Zuordnung nur geraten.
    person_domain = _norm(person.get("firmen_domain"))
    if person_domain and person_domain != _norm(zeile["domain"]):
        return "company_mismatch"
    return "valid"


def _guete_grund(zeile: dict, person: dict) -> str:
    if zeile["company_match_status"] == "no_match":
        return "Person gehört nachweislich zu einer anderen Firma"
    gruende = []
    if zeile["company_match_status"] != "strong_match":
        gruende.append(f"Firmen-Zuordnung nur {zeile['company_match_status']}")
    if not (person.get("titel") or "").strip():
        kopf = (person.get("headline") or "")[:60]
        gruende.append(f"kein Stellentitel im Profil"
                       + (f", nur headline '{kopf}'" if kopf else ""))
    elif person.get("rang") is None or person["rang"] >= len(RANG_GRUPPEN):
        gruende.append(f"Titel '{person.get('titel')}' ist kein "
                       f"Entscheider-Titel")
    return "; ".join(gruende) or "Entscheider nicht belastbar"


def _eine_firma(quelle, firma, zeile, roh_speichern, such_credits=None):
    """Firma -> Person -> Kontaktdaten. Fuellt zeile direkt."""
    from pipeline.sources.fullenrich import PREIS_SUCHTREFFER

    if such_credits is None:
        such_credits = [0.0]
    domain = zeile["domain"]
    if not domain:
        zeile["api_status"] = "skipped_no_domain"
        zeile["error_message"] = "Firma hat keine Domain"
        return

    fe_firma = quelle.firma_finden(domain)
    zeile["fullenrich_company_found"] = 1 if fe_firma else 0
    if fe_firma:
        such_credits[0] += PREIS_SUCHTREFFER
        zeile["fullenrich_company_id"] = str(fe_firma.get("id") or "")
        zeile["fullenrich_company_name"] = fe_firma.get("name") or ""
        zeile["fullenrich_company_domain"] = _domain(
            fe_firma.get("domain") or fe_firma.get("website"))

    personen = quelle.personen_suchen(domain, TITEL_VARIANTEN)
    such_credits[0] += PREIS_SUCHTREFFER * len(personen)
    gewaehlt, alternativen = entscheider_waehlen(personen, firma)

    status, punkte, belege = firmen_abgleich(firma, fe_firma, gewaehlt)
    zeile["company_match_status"] = status
    zeile["company_match_score"] = punkte
    zeile["company_match_evidence"] = belege

    if not gewaehlt:
        zeile["api_status"] = "no_decision_maker"
        zeile["decision_maker_status"] = "none"
        return

    zeile["decision_maker_found"] = 1
    zeile["decision_maker_first_name"] = gewaehlt["first_name"]
    zeile["decision_maker_last_name"] = gewaehlt["last_name"]
    zeile["decision_maker_name"] = gewaehlt["name"]
    zeile["decision_maker_title"] = gewaehlt["titel"]
    zeile["decision_maker_headline"] = gewaehlt["headline"]
    zeile["decision_maker_linkedin"] = gewaehlt["linkedin"]
    zeile["decision_maker_rank"] = gewaehlt["rang"]
    zeile["decision_maker_reason"] = gewaehlt.get("rang_grund", "")
    zeile["decision_maker_title_source"] = gewaehlt.get("titel_quelle", "")
    zeile["decision_maker_headline_evidence"] = gewaehlt.get(
        "headline_beleg", "")
    zeile["alternative_candidates"] = json.dumps(
        [{"name": a["name"], "titel": a["titel"], "rang": a["rang"]}
         for a in alternativen], ensure_ascii=False)

    # ---- KOSTEN-TOR: hier wird entschieden, ob die teuren Credits
    # ueberhaupt ausgegeben werden. Eine Mobilnummer kostet 10 Credits -
    # die darf NICHT fuer eine Person fliessen, von der wir nicht wissen,
    # ob sie ueberhaupt entscheidet.
    zeile["decision_maker_status"] = _entscheider_guete(zeile, gewaehlt)
    if zeile["decision_maker_status"] != "valid":
        zeile["api_status"] = f"skipped_{zeile['decision_maker_status']}"
        zeile["manual_review_reason"] = _guete_grund(zeile, gewaehlt)
        return

    enrichment_id = quelle.anreicherung_starten([{
        "first_name": gewaehlt["first_name"] or gewaehlt["name"].split(" ")[0],
        "last_name": gewaehlt["last_name"]
        or " ".join(gewaehlt["name"].split(" ")[1:]),
        "domain": domain,
        "company_name": firma.get("name", ""),
        "linkedin_url": gewaehlt["linkedin"] or None,
        "custom": {"company_id": str(firma.get("id") or "")},
    }], name=f"FE {firma.get('name','')[:60]}")
    if not enrichment_id:
        zeile["api_status"] = "enrich_not_started"
        return

    ergebnis = quelle.anreicherung_abholen(enrichment_id)
    zeile["api_status"] = str(ergebnis.get("status") or "UNKNOWN")
    zeile["credits_used"] = int(
        (ergebnis.get("cost") or {}).get("credits") or 0)
    if roh_speichern:
        zeile["raw_response"] = _saubern(ergebnis)

    daten = ergebnis.get("data") or []
    if not daten:
        return
    kontakt = (daten[0] or {}).get("contact_info") or {}

    # Drei Kategorien, streng getrennt (Auftrag 24.08.2026):
    #   individual_work_email   j.smith@firma.de   <- DAS will Oliver
    #   generic_company_email   info@firma.de      <- erfuellt es NICHT
    #   private_personal_email  j.smith@gmail.com  <- getrennt gespeichert
    arbeit = beste_mail(kontakt.get("work_emails"))
    if arbeit:
        zeile["work_email"] = arbeit["email"]
        zeile["work_email_status"] = arbeit.get("status", "")
        generisch = ist_sammeladresse(arbeit["email"])
        zeile["work_email_is_generic"] = 1 if generisch else 0
        if generisch:
            zeile["generic_company_email"] = arbeit["email"]
        else:
            zeile["individual_work_email"] = arbeit["email"]
    persoenlich = beste_mail(kontakt.get("personal_emails"))
    if persoenlich:
        zeile["personal_email"] = persoenlich["email"]
        zeile["personal_email_status"] = persoenlich.get("status", "")
        zeile["private_personal_email"] = persoenlich["email"]

    _telefone_einordnen(kontakt.get("phones"), zeile)


def _telefone_einordnen(telefone, zeile):
    """Mobil / Durchwahl / Firmenzentrale trennen.

    Die Zentrale darf NIE als Durchwahl oder Mobilnummer zaehlen - das
    ist ausdruecklich Olivers Anforderung. Verglichen wird gegen die
    Firmennummer aus unserer eigenen Datenbank."""
    eintraege = [t for t in (telefone or [])
                 if isinstance(t, dict) and t.get("number")]
    zeile["phones_raw"] = json.dumps(eintraege, ensure_ascii=False)
    zentrale = _ziffern(zeile.get("company_phone"))

    unbekannte = []
    for eintrag in eintraege:
        nummer = eintrag["number"]
        if zentrale and _ziffern(nummer) == zentrale:
            # Identisch mit der Firmenzentrale - keine persoenliche Nummer.
            if not zeile["direct_phone_status"]:
                zeile["direct_phone_status"] = "company_switchboard_ignored"
            continue
        art = telefon_art(nummer)
        if art == "mobil" and not zeile["mobile_phone"]:
            zeile["mobile_phone"] = nummer
            zeile["mobile_phone_status"] = "mobile_by_prefix"
        elif art == "festnetz" and not zeile["direct_phone"]:
            zeile["direct_phone"] = nummer
            zeile["direct_phone_status"] = "landline_by_prefix"
        elif art == "unbekannt":
            unbekannte.append(nummer)

    # Nicht einzuordnende Nummern (Ausland, unklare Schreibweise) werden
    # NICHT als Mobil oder Durchwahl ausgegeben - das waere geraten.
    # Sie gehen aber auch nicht verloren.
    if unbekannte:
        zeile["unknown_phones"] = json.dumps(unbekannte, ensure_ascii=False)
        if not (zeile["mobile_phone"] or zeile["direct_phone"]):
            zeile["direct_phone_status"] = "type_unknown_not_counted"
