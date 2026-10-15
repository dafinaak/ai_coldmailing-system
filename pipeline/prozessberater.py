"""Process consultants in postal regions 30-39: search area and Excel list.

A different target group than the IT lists (Dafina, 14.10.2026): firms
that sell process, organisation and REFA consulting. The IT chain
(branchen_filter, automation check, zone folders) is built to throw
exactly these firms out, so this path stands on its own: Google Maps
finds them, the postal code decides whether they belong, and the result
is a plain firm list in Excel. No profile filter, no Dropcontact - a
person with an email is a later, separate step.

The search area is NOT the circle of a postal region. Regions 30, 31, 35
and 36 are scattered, so their circles end up at the 120 km cap and
overlap each other - Apify bills every place found, and a place found
by two regions is paid twice. Instead each region becomes a MultiPolygon
of grid cells, one cell for every cell that holds a postal code centre
of that region. The postal code filter afterwards decides membership.

The CLI with the paid call is werkzeuge/prozessberater-maps.py.
"""
from __future__ import annotations

import re
from collections import defaultdict
from urllib.parse import urlparse

from pipeline.plz_geo import plz_tabelle

# The terms exactly as Dafina sent them on 14.10.2026, duplicates removed
# ("Prozessoptimierung" was listed twice).
SEARCH_TERMS = [
    "Prozessberater",
    "Prozess Consultant",
    "Prozess Consulting",
    "Prozessoptimierer",
    "Prozessspezialist",
    "Prozessexperte",
    "Prozessoptimierung",
    "Prozessberatung",
    "Prozessmanagement",
    "Prozess- und Organisationsberater",
    "Prozess- und Organisationsberatung",
    "Prozess- und Strategieberatung",
    "Organisations- und Prozessberatung",
    "Geschäftsprozessanalyse",
    "Geschäftsberatung",
    "Geschäftsberater",
    "Consulting im Prozessmanagement",
    "Prozessberatung für Geschäftsprozesse",
    "REFA Consulting",
    "REFA Berater",
    "REFA Beratung",
]

REGIONS = [str(n) for n in range(30, 40)]

# Grid cell in degrees: about 17 x 17 km at latitude 52.
CELL_LAT = 0.15
CELL_LON = 0.25

# A main category with none of these word stems is marked for a human
# look - "Geschäftsberatung" also brings tax advisors, insurers and the
# like. Only a hint in its own column; nothing is removed for it.
CONSULTING_STEMS = ("berat", "consult", "prozess", "refa", "organisation",
                    "management", "coach", "schulung", "training")
# "Steuerberater" also contains "berat" - these stems flag a category
# even when a consulting stem is in it.
OFF_TARGET_STEMS = ("steuer", "finanz", "versicher", "immobil", "anwalt",
                    "makler", "bank", "kredit", "marketing", "werbe",
                    "energie", "it-", "computer", "software")


def category_needs_check(category: str) -> bool:
    text = category.lower()
    return (any(s in text for s in OFF_TARGET_STEMS)
            or not any(s in text for s in CONSULTING_STEMS))


def region_cells(region: str, table=None) -> list[tuple[int, int]]:
    """Grid cells (row, column) holding a postal code centre of the region."""
    table = table if table is not None else plz_tabelle()
    cells = set()
    for plz, entry in table.items():
        if str(plz).startswith(region):
            cells.add((int(entry["lat"] // CELL_LAT),
                       int(entry["lon"] // CELL_LON)))
    return sorted(cells)


def cells_geojson(cells) -> dict:
    """The cells as one GeoJSON MultiPolygon (Apify: customGeolocation)."""
    polygons = []
    for row, col in cells:
        s, w = row * CELL_LAT, col * CELL_LON
        n, e = s + CELL_LAT, w + CELL_LON
        ring = [[round(w, 6), round(s, 6)], [round(e, 6), round(s, 6)],
                [round(e, 6), round(n, 6)], [round(w, 6), round(n, 6)],
                [round(w, 6), round(s, 6)]]
        polygons.append([ring])
    return {"type": "MultiPolygon", "coordinates": polygons}


def max_cost(regions: int, terms: int, cap: int, price: float) -> float:
    """Upper bound: every term hits its cap in every region."""
    return regions * terms * cap * price


def _domain(website: str) -> str:
    if not website:
        return ""
    netloc = urlparse(website).netloc.lower() or website.split("/")[0].lower()
    return netloc[4:] if netloc.startswith("www.") else netloc


def _plz(item: dict) -> str:
    plz = str(item.get("postalCode") or "").strip()
    if not plz:
        # Some places only carry the postal code inside the address line.
        found = re.search(r"\b(\d{5})\b", str(item.get("address") or ""))
        plz = found.group(1) if found else ""
    return plz


def build_rows(datasets: dict) -> tuple[list[dict], dict]:
    """Merge the raw Apify items of all regions into one firm list.

    datasets: {region: [apify items]}. Kept is a place whose postal code
    starts with 30-39; a place without a postal code is left out, because
    then nothing proves it sits in the area (same rule as the zone lists,
    Dafina 10.09.2026). One row per Google place; a website seen on
    several places stays one row, with the other locations counted.
    """
    places: dict = {}
    terms_of = defaultdict(set)
    counts = {"geliefert": 0, "ohne_plz": 0, "ausserhalb_30_39": 0}
    for region, items in datasets.items():
        for item in items:
            counts["geliefert"] += 1
            key = (item.get("placeId") or item.get("url")
                   or f"{item.get('title')}|{item.get('address')}")
            if item.get("searchString"):
                terms_of[key].add(item["searchString"])
            if key in places:
                continue
            plz = _plz(item)
            if not plz:
                counts["ohne_plz"] += 1
                places[key] = None
                continue
            if plz[:2] not in REGIONS:
                counts["ausserhalb_30_39"] += 1
                places[key] = None
                continue
            places[key] = (region, plz, item)

    rows, by_domain = [], {}
    for key, value in places.items():
        if value is None:
            continue
        region, plz, item = value
        website = item.get("website") or ""
        domain = _domain(website)
        if domain and domain in by_domain:
            by_domain[domain]["Weitere Standorte"] += 1
            by_domain[domain]["Suchbegriffe"] |= terms_of[key]
            continue
        categories = item.get("categories") or (
            [item["categoryName"]] if item.get("categoryName") else [])
        main = item.get("categoryName") or (categories[0] if categories else "")
        closed = ("dauerhaft geschlossen" if item.get("permanentlyClosed")
                  else "voruebergehend geschlossen"
                  if item.get("temporarilyClosed") else "")
        row = {
            "Firma": item.get("title") or "",
            "Kategorie": main,
            "Kategorie pruefen": "ja" if category_needs_check(main) else "",
            "Weitere Kategorien": ", ".join(c for c in categories if c != main),
            "Webseite": website,
            "Telefon": item.get("phone") or "",
            "Strasse": item.get("street") or "",
            "PLZ": plz,
            "Ort": item.get("city") or "",
            "PLZ-Region": plz[:2],
            "Bewertung": item.get("totalScore") or "",
            "Anzahl Bewertungen": item.get("reviewsCount") or "",
            "Status": closed,
            "Google Maps": item.get("url") or "",
            "Suchbegriffe": set(terms_of[key]),
            "Weitere Standorte": 0,
        }
        rows.append(row)
        if domain:
            by_domain[domain] = row
    for row in rows:
        row["Suchbegriffe"] = ", ".join(sorted(row["Suchbegriffe"]))
    rows.sort(key=lambda r: (r["PLZ"], r["Firma"].lower()))
    counts["firmen"] = len(rows)
    return rows, counts


COLUMNS = [("Nr", 6), ("Firma", 40), ("Kategorie", 28),
           ("Kategorie pruefen", 10), ("Weitere Kategorien", 34),
           ("Webseite", 34), ("Telefon", 18), ("Strasse", 28), ("PLZ", 8),
           ("Ort", 20), ("PLZ-Region", 8), ("Bewertung", 9),
           ("Anzahl Bewertungen", 9), ("Status", 14),
           ("Weitere Standorte", 9), ("Suchbegriffe", 40),
           ("Google Maps", 40)]


def write_excel(path, rows: list[dict], counts: dict,
                per_term: list[tuple]) -> None:
    """Sheet "Firmen" plus sheet "Bericht" (counts, per region and term).

    per_term: (region, term, places delivered, cap) - a term that hit its
    cap may have more firms than we got, so the report says so.
    """
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = "Firmen"
    sheet.append([name for name, _ in COLUMNS])
    for nr, row in enumerate(rows, 1):
        sheet.append([nr if name == "Nr" else row.get(name, "")
                      for name, _ in COLUMNS])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column, (_, width) in enumerate(COLUMNS, 1):
        sheet.column_dimensions[get_column_letter(column)].width = width

    report = wb.create_sheet("Bericht")
    report.append(["Gesamt", "Anzahl"])
    for name, value in [
            ("Orte von Google Maps geliefert", counts["geliefert"]),
            ("ohne PLZ (nicht aufgenommen)", counts["ohne_plz"]),
            ("PLZ ausserhalb 30-39 (nicht aufgenommen)",
             counts["ausserhalb_30_39"]),
            ("Firmen in der Liste", counts["firmen"])]:
        report.append([name, value])
    report.append([])
    report.append(["PLZ-Region", "Suchbegriff", "Orte geliefert",
                   "Grenze erreicht"])
    for region, term, delivered, cap in per_term:
        report.append([region, term, delivered,
                       "ja - es gibt wohl mehr" if delivered >= cap else ""])
    for cell in report[1]:
        cell.font = Font(bold=True)
    report.column_dimensions["A"].width = 42
    report.column_dimensions["B"].width = 40
    report.column_dimensions["C"].width = 14
    report.column_dimensions["D"].width = 24
    wb.save(path)
