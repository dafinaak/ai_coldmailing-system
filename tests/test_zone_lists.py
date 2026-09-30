"""All contacts of ten zones in one Excel sheet (Dafina, 29.09.2026).

"Me i bo ne nje list te vetme te gjitha kontaktet prej 40 deri 49 ne nje
excel, pastaj 50 deri 59 tjetrin" - and after seeing one sheet per zone:
"boni kejt te nje sheet mos i ndaj hiq". So one file per ten zones with one
sheet: the contacts of every zone under each other, columns exactly as on
the zone lists, one last column "Zona", "Nr" counting through. Only what can
be sent goes in - the catch-all addresses and the check sheet stay on the
zone lists, which themselves stay untouched.
"""
import openpyxl
import pytest

from pipeline.zone_lists import merge_zone_lists, newest_per_zone

HEADER = ["Nr", "Firma", "Person", "E-Mail"]


def _zone_list(path, rows, extra=None, header=HEADER):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Versandfertig"
    ws.append(header)
    for row in rows:
        ws.append(row)
    for name, (sheet_header, sheet_rows) in (extra or {}).items():
        sheet = wb.create_sheet(name)
        sheet.append(sheet_header)
        for row in sheet_rows:
            sheet.append(row)
    wb.save(path)
    return path


def _rows(path, sheet="Versandfertig"):
    return list(openpyxl.load_workbook(path)[sheet].iter_rows(values_only=True))


def test_all_zones_in_one_sheet_counted_through(tmp_path):
    a = _zone_list(tmp_path / "a.xlsx", [[1, "A GmbH", "Anna Alt", "anna@a.de"],
                                         [2, "B GmbH", "Bernd Bau", "bernd@b.de"]])
    b = _zone_list(tmp_path / "b.xlsx", [[1, "C GmbH", "Clara Cent", "clara@c.de"]])

    merged = merge_zone_lists({"41": b, "40": a}, tmp_path / "out.xlsx")

    assert openpyxl.load_workbook(merged).sheetnames == ["Versandfertig"]
    assert _rows(merged) == [
        ("Nr", "Firma", "Person", "E-Mail", "Zona"),
        (1, "A GmbH", "Anna Alt", "anna@a.de", "40"),
        (2, "B GmbH", "Bernd Bau", "bernd@b.de", "40"),
        (3, "C GmbH", "Clara Cent", "clara@c.de", "41"),
    ]


def test_addresses_that_cannot_be_sent_stay_out(tmp_path):
    # A catch-all address next to the sendable ones could be sent by mistake.
    extra = {"Bezahlt, nicht versandfähig": (["Nr", "Firma", "E-Mail"],
                                             [[1, "D GmbH", "x@d.de"]]),
             "Zur Kontrolle": (["Firma", "Person", "Warum"],
                               [["A GmbH", "Anna Alt", "Herr laut Dropcontact"]])}
    a = _zone_list(tmp_path / "a.xlsx", [[1, "A GmbH", "Anna Alt", "anna@a.de"]],
                   extra=extra)

    merged = merge_zone_lists({"40": a}, tmp_path / "out.xlsx")

    assert openpyxl.load_workbook(merged).sheetnames == ["Versandfertig"]
    assert "x@d.de" not in str(_rows(merged))


def test_the_newest_list_of_each_zone_in_the_range_is_taken(tmp_path):
    for name in ("IT-Liste-Emails-Zona40-FERTIG-20260929-1101.xlsx",
                 "IT-Liste-Emails-Zona40-FERTIG-20260929-1308.xlsx",
                 "IT-Liste-Emails-Zona41-FERTIG-20260929-1135.xlsx",
                 "IT-Liste-Emails-Zona39-FERTIG-20260910-1121.xlsx",
                 "IT-Liste-Emails-Zona52-FERTIG-20261010-0900.xlsx",
                 "IT-Liste-Emails-Zona41-20260929-1134.xlsx",
                 "IT-Liste-Emails-Zonat-40-49-20260929-1400.xlsx"):
        _zone_list(tmp_path / name, [])

    found = newest_per_zone(tmp_path, 40)

    assert {zone: path.name for zone, path in found.items()} == {
        "40": "IT-Liste-Emails-Zona40-FERTIG-20260929-1308.xlsx",
        "41": "IT-Liste-Emails-Zona41-FERTIG-20260929-1135.xlsx"}


def test_other_columns_stop_the_merge(tmp_path):
    # Rows under a header they do not belong to would put e-mails into the
    # wrong column - better no file than a wrong one.
    a = _zone_list(tmp_path / "a.xlsx", [[1, "A GmbH", "Anna Alt", "anna@a.de"]])
    b = _zone_list(tmp_path / "b.xlsx", [[1, "C GmbH", "clara@c.de"]],
                   header=["Nr", "Firma", "E-Mail"])

    with pytest.raises(ValueError, match="41"):
        merge_zone_lists({"40": a, "41": b}, tmp_path / "out.xlsx")
    assert not (tmp_path / "out.xlsx").exists()


def test_one_person_on_two_zone_lists_stops_the_merge(tmp_path):
    # The zone lists drop such a person already (the lower zone keeps it).
    # If one slipped through, one campaign would write to them twice.
    a = _zone_list(tmp_path / "a.xlsx", [[1, "A GmbH", "Anna Alt", "anna@a.de"]])
    b = _zone_list(tmp_path / "b.xlsx", [[1, "A2 GmbH", "Anna Alt", "anna@a.de"]])

    with pytest.raises(ValueError, match="anna@a.de"):
        merge_zone_lists({"40": a, "41": b}, tmp_path / "out.xlsx")
