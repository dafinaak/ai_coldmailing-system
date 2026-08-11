"""Turn a checked Excel dispatch list back into the JSON the run reads.

The company records with website, domain and manager names live in the
JSON that the lead sourcing produced. The Excel file is what humans
check and mark. This script keeps the two in step: it takes the final
Excel list and writes out only those companies, with their full JSON
record plus the salutation column that was added by hand.

A company in the Excel file that has no match in the JSON is reported
rather than silently dropped - a missing record would quietly shrink
the run.

Usage:
    python -m pipeline.liste_als_json <liste.xlsx> <quelle.json> <ziel.json>
"""

from __future__ import annotations

import json
import sys
import unicodedata
from pathlib import Path

import openpyxl


def schluessel(name: object) -> str:
    text = unicodedata.normalize("NFKD", str(name or ""))
    return "".join(z for z in text if z.isprintable()).strip().lower()


def main(liste: Path, quelle: Path, ziel: Path) -> None:
    ws = openpyxl.load_workbook(liste).worksheets[0]
    kopf = [c.value for c in ws[1]]
    spalte_anrede = kopf.index("Anrede")

    gewuenscht = {}
    doppelt: dict[str, list[str]] = {}
    for zeile in ws.iter_rows(min_row=2, values_only=True):
        if zeile[1] is None:
            continue
        name = schluessel(zeile[1])
        if name in gewuenscht:
            doppelt.setdefault(name, [str(zeile[1])]).append(str(zeile[6]))
        gewuenscht[name] = zeile[spalte_anrede]

    records = json.loads(quelle.read_text(encoding="utf-8"))
    nach_name = {schluessel(r.get("name")): r for r in records}

    raus, fehlend = [], []
    for name, anrede in gewuenscht.items():
        record = nach_name.get(name)
        if record is None:
            fehlend.append(name)
            continue
        raus.append({**record, "anrede": anrede})

    ziel.write_text(
        json.dumps(raus, ensure_ascii=False, indent=1), encoding="utf-8"
    )
    print(f"Excel-Liste: {ws.max_row - 1} Zeilen -> {len(gewuenscht)} Firmennamen")
    if doppelt:
        print(f"ACHTUNG - {len(doppelt)} Firma(en) stehen mehrfach in der Liste.")
        print("Es wird nur EIN Datensatz je Firma uebernommen; welche Webseite")
        print("die richtige ist, gehoert von Hand entschieden:")
        for name, eintraege in doppelt.items():
            record = nach_name.get(name, {})
            print(f"     - {eintraege[0]} -> genommen: {record.get('website')}")
    print(f"geschrieben: {len(raus)} Firmen -> {ziel}")
    if fehlend:
        print(f"OHNE Datensatz in {quelle.name} ({len(fehlend)}):")
        for name in fehlend:
            print(f"     - {name}")


if __name__ == "__main__":
    if len(sys.argv) != 4:
        raise SystemExit(__doc__)
    main(Path(sys.argv[1]), Path(sys.argv[2]), Path(sys.argv[3]))
