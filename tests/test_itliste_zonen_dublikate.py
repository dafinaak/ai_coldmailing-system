"""Ein Mensch, eine Mail - auch wenn nur noch Sammellisten herumliegen.

Die dritte Porta in werkzeuge/zona32-itliste-final.py haelt jemanden aus
der Liste, der schon in einer Zone mit KLEINERER Nummer steht: eine Firma
mit zwei Standorten faellt sonst in zwei Listen, und ihr Entscheider
bekaeme zwei Mails aus derselben Kampagne (der Fehler vom 17.08.2026).

Dafuer las die Porta die Einzel-Listen je Zone. Am 07.10.2026 wurden die
auf Dafinas Wunsch geloescht - pro Dekade bleibt nur noch EINE
Sammelliste. Danach fand der glob nichts mehr, die Menge blieb leer, und
die Porta liess stillschweigend alles durch: in der Dekade 60-69 standen
anschliessend sieben Menschen, die schon in 30-39 bzw. 40-49 standen
(IT-HAUS, Ratiodata, Concat, Medialine, H&G, C.B.C., Compose IT).

Deshalb liest die Porta jetzt beides - und ein Lauf, der ueberhaupt keine
Vergleichsliste findet, obwohl es kleinere Zonen gibt, sagt das laut.
"""
import importlib.util
from pathlib import Path

import openpyxl

PROJEKT = Path(__file__).resolve().parent.parent


def _werkzeug(projekt: Path):
    pfad = PROJEKT / "werkzeuge" / "zona32-itliste-final.py"
    spec = importlib.util.spec_from_file_location("itliste_final", pfad)
    modul = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(modul)
    modul.PROJEKT = projekt
    return modul


def _liste(pfad: Path, person="Anna Alt", email="a@firma.de", plz="40210"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Nr", "Firma", "Person", "Position", "E-Mail", "Anrede",
               "Hinweis", "Telefon (Person)", "Telefon (Firma)", "Webseite",
               "PLZ", "Ort"])
    ws.append([1, "Firma GmbH", person, "", email, "Frau Alt", "", "", "",
               "https://firma.de", plz, "Stadt"])
    wb.save(pfad)


def test_einzel_listen_werden_gelesen(tmp_path):
    """Der alte Weg bleibt."""
    _liste(tmp_path / "IT-Liste-Emails-Zona40-FERTIG-20261007-1000.xlsx")
    modul = _werkzeug(tmp_path)
    schon = modul.kontakte_kleinerer_zonen("60")
    assert ("email", "a@firma.de") in schon
    assert ("person", "anna alt") in schon


def test_sammellisten_werden_gelesen_wenn_einzel_fehlen(tmp_path):
    """Der Fall vom 07.10.2026."""
    _liste(tmp_path / "IT-Liste-Emails-Zonat-30-39-20261007-1108.xlsx",
           person="Bernd Bach", email="b@haus.de")
    modul = _werkzeug(tmp_path)
    schon = modul.kontakte_kleinerer_zonen("60")
    assert ("email", "b@haus.de") in schon
    assert ("person", "bernd bach") in schon


def test_in_der_eigenen_dekade_zaehlt_die_zeile_ihre_eigene_plz(tmp_path):
    """Die Sammelliste 60-69 enthaelt kleinere Zonen UND die eigene.

    Darum entscheidet die Postleitzahl der Zeile, nicht der Dateiname:
    fuer Zone 69 ist eine 60er-Zeile eine kleinere Zone, eine 69er-Zeile
    aber die eigene - und die darf sich nicht selbst wegstreichen.
    """
    pfad = tmp_path / "IT-Liste-Emails-Zonat-60-69-20261007-1438.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Nr", "Firma", "Person", "Position", "E-Mail", "Anrede",
               "Hinweis", "Telefon (Person)", "Telefon (Firma)", "Webseite",
               "PLZ", "Ort"])
    ws.append([1, "A GmbH", "Carla Cle", "", "c@ff.de", "Frau Cle", "", "",
               "", "https://ff.de", "60311", "Frankfurt"])
    ws.append([2, "B GmbH", "Dirk Dorn", "", "d@gg.de", "Herr Dorn", "", "",
               "", "https://gg.de", "69115", "Heidelberg"])
    wb.save(pfad)

    modul = _werkzeug(tmp_path)
    schon = modul.kontakte_kleinerer_zonen("69")
    assert ("email", "c@ff.de") in schon        # zona 60 < 69
    assert ("email", "d@gg.de") not in schon    # zona 69 = vetja
    assert modul.kontakte_kleinerer_zonen("60") == set()


def test_erste_zone_hat_nichts_zu_vergleichen(tmp_path):
    modul = _werkzeug(tmp_path)
    assert modul.kontakte_kleinerer_zonen("32") == set()


def test_dekade_vergleicht_die_zone_der_zeile_nicht_zwei_stellen(tmp_path):
    """In der Sammelliste entscheidet die Zone der Zeile, nicht ihre
    ersten zwei Stellen.

    Seit dem 09.10.2026 liegen mehrere Staedte in derselben Dekade und
    heissen darum drei Stellen lang: Hagen "58", Luedenscheid "585",
    Iserlohn "586", Arnsberg "59", Hamm "590". Zwei Stellen koennen die
    nicht auseinanderhalten: fuer Hamm ("590") ergibt Arnsberg (59755)
    "597" < "590" = falsch, also galt die kleinere Zone als groesser.

    Folge: eine Person, die in Arnsberg und in Hamm sitzt, bleibt auf
    beiden Listen - und bekaeme dieselbe Kampagne zweimal, genau der
    Fehler vom 17.08.2026.
    """
    pfad = tmp_path / "IT-Liste-Emails-Zonat-50-59-20261010-0900.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Nr", "Firma", "Person", "Position", "E-Mail", "Anrede",
               "Hinweis", "Telefon (Person)", "Telefon (Firma)", "Webseite",
               "PLZ", "Ort"])
    for nr, (email, plz, ort) in enumerate([
            ("hagen@f.de", "58089", "Hagen"),
            ("luedenscheid@f.de", "58507", "Lüdenscheid"),
            ("iserlohn@f.de", "58636", "Iserlohn"),
            ("arnsberg@f.de", "59755", "Arnsberg"),
            ("hamm@f.de", "59063", "Hamm")], 1):
        ws.append([nr, "F GmbH", f"P {nr}", "", email, "", "", "", "",
                   "https://f.de", plz, ort])
    wb.save(pfad)

    modul = _werkzeug(tmp_path)
    modul.ZONEN = {"58": {}, "585": {}, "586": {}, "59": {}, "590": {}}

    schon = modul.kontakte_kleinerer_zonen("590")
    assert ("email", "arnsberg@f.de") in schon     # zona 59 < 590
    assert ("email", "hagen@f.de") in schon        # zona 58 < 590
    assert ("email", "iserlohn@f.de") in schon     # zona 586 < 590
    assert ("email", "hamm@f.de") not in schon     # zona 590 = vetja

    # dhe anasjelltas: per Arnsbergun, Hamm-i eshte zone me e madhe
    schon = modul.kontakte_kleinerer_zonen("59")
    assert ("email", "hamm@f.de") not in schon
    assert ("email", "arnsberg@f.de") not in schon
    assert ("email", "hagen@f.de") in schon
