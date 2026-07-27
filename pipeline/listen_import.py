"""Import einer vorgefertigten Lead-Liste (xlsx) in das Firmen-Format der
Pipeline (Bauplan 2026-07-27, Baustein 2).

Erwartete Spalten (wie in "IT-Dienstleister PLR 30-39 - 319 HQ-only Leads"):
ID | Firma | Kategorie | PLZ | Standort | Website | Telefon | E-Mail |
Geschäftsführer/Entscheider | Mitarbeiter-Hinweis | Quelle

Wichtig (Stichproben-Fund 27.07.2026): Der Geschaeftsfuehrer-Name aus der
Liste ist ein HINWEIS, kein Fakt - von 8 gepruefte Zeilen stimmten nur 5;
einmal stand sogar ein Firmenname im Personen-Feld. Er wird deshalb als
"gf_name_liste" mitgefuehrt und von der Impressum-Stufe gegen die echte
Webseite geprueft, nie ungeprueft zu Dropcontact gegeben.
"""
from pathlib import Path
from urllib.parse import urlparse


def _s(wert) -> str:
    return str(wert).strip() if wert is not None else ""


def _domain(url: str) -> str:
    netloc = urlparse(_s(url)).netloc.lower()
    if not netloc:
        netloc = _s(url).lower().split("/")[0]
    return netloc[4:] if netloc.startswith("www.") else netloc


def liste_lesen(xlsx_pfad, region_praefixe: tuple = ("3",)) -> list:
    """Liest die xlsx-Liste und gibt Firmen-Dicts im Pipeline-Format zurueck.
    Zusatzfelder gegenueber dem Apify-Format: plz, telefon, gf_name_liste,
    vorhandene_email, quelle, ausserhalb_region (PLZ passt nicht zu
    region_praefixe oder fehlt - Olivers "nicht vergessen"-Regel)."""
    import openpyxl
    wb = openpyxl.load_workbook(Path(xlsx_pfad), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    zeilen = list(ws.iter_rows(values_only=True))
    firmen = []
    for r in zeilen[1:]:
        if not any(r):
            continue
        website = _s(r[5])
        firmen.append({
            "name": _s(r[1]), "website": website, "domain": _domain(website),
            "address": (_s(r[3]) + " " + _s(r[4])).strip(),
            "categories": [_s(r[2])] if _s(r[2]) else [],
            "plz": _s(r[3]), "telefon": _s(r[6]),
            "vorhandene_email": _s(r[7]), "gf_name_liste": _s(r[8]),
            "quelle": _s(r[10]) if len(r) > 10 else "",
            "ausserhalb_region": not _s(r[3]).startswith(tuple(region_praefixe)),
        })
    return firmen
