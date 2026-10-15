from openpyxl import load_workbook

from pipeline import prozessberater as pb

TABLE = {
    "30159": {"ort": "Hannover", "lat": 52.37, "lon": 9.73},
    "30161": {"ort": "Hannover", "lat": 52.38, "lon": 9.72},
    "38100": {"ort": "Braunschweig", "lat": 52.26, "lon": 10.52},
    "40210": {"ort": "Duesseldorf", "lat": 51.22, "lon": 6.79},
}


def test_cells_only_for_the_region():
    cells = pb.region_cells("30", TABLE)
    assert len(cells) == 1               # both Hannover codes share a cell
    geo = pb.cells_geojson(cells)
    assert geo["type"] == "MultiPolygon"
    ring = geo["coordinates"][0][0]
    assert ring[0] == ring[-1]
    lons = [p[0] for p in ring]
    lats = [p[1] for p in ring]
    assert min(lons) <= 9.73 <= max(lons) and min(lats) <= 52.37 <= max(lats)


def test_terms_have_no_duplicates():
    assert len(pb.SEARCH_TERMS) == len(set(pb.SEARCH_TERMS)) == 21


def _item(pid, title, plz, website="", term="Prozessberatung",
          category="Unternehmensberater"):
    return {"placeId": pid, "title": title, "postalCode": plz,
            "website": website, "searchString": term,
            "categoryName": category, "city": "X"}


def test_rows_keep_only_30_39_with_postal_code():
    rows, counts = pb.build_rows({
        "30": [_item("a", "Alpha", "30159"),
               _item("b", "Beta", ""),            # no postal code
               _item("c", "Gamma", "29221")],     # outside 30-39
        "31": [_item("a", "Alpha", "30159", term="REFA Berater")],
    })
    assert [r["Firma"] for r in rows] == ["Alpha"]
    assert rows[0]["Suchbegriffe"] == "Prozessberatung, REFA Berater"
    assert counts == {"geliefert": 4, "ohne_plz": 1,
                      "ausserhalb_30_39": 1, "firmen": 1}


def test_same_website_is_one_row_and_odd_category_is_flagged():
    rows, _ = pb.build_rows({"30": [
        _item("a", "Alpha Hannover", "30159", "https://www.alpha.de"),
        _item("b", "Alpha Celle", "31234", "https://alpha.de/celle"),
        _item("c", "Steuer GmbH", "30161", category="Steuerberater"),
    ]})
    alpha = [r for r in rows if r["Firma"].startswith("Alpha")]
    assert len(alpha) == 1 and alpha[0]["Weitere Standorte"] == 1
    steuer = [r for r in rows if r["Firma"] == "Steuer GmbH"][0]
    assert steuer["Kategorie pruefen"] == "ja"
    assert alpha[0]["Kategorie pruefen"] == ""


def test_excel_has_list_and_report(tmp_path):
    rows, counts = pb.build_rows({"33": [_item("a", "Alpha", "33602")]})
    path = tmp_path / "liste.xlsx"
    pb.write_excel(path, rows, counts, [("33", "REFA Berater", 60, 60)])
    wb = load_workbook(path)
    assert wb["Firmen"]["B2"].value == "Alpha"
    report = [r for r in wb["Bericht"].iter_rows(values_only=True)]
    assert ("33", "REFA Berater", 60, "ja - es gibt wohl mehr") in report


def test_category_check():
    assert not pb.category_needs_check("Unternehmensberater")
    assert not pb.category_needs_check("Consultant")
    assert pb.category_needs_check("Steuerberater")
    assert pb.category_needs_check("IT-Berater")
    assert pb.category_needs_check("Bäckerei")
