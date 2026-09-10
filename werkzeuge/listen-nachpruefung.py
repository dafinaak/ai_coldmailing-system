#!/usr/bin/env python3
"""Rigjykimi i profilit IT me rregullin e SOTEM, per listat e vjetra.

Pse: rregulli i 31.08.2026 (AGENTS.md) thote qe firmat me softuer te
vet, me thelb sigurie, partneret e produkteve te huaja dhe agjencite
jane JASHTE - edhe kur ofrojne mirembajtje IT. Dhe thote shprehimisht qe
rregulli vlen edhe per listat e gatshme. Por zonat 32, 33 dhe 34 u gjykuan
ne gusht me rregullin e vjeter (plus hapin "shansi i dyte", qe i kthente
brenda pikerisht keta), dhe askush s'i rishikoi: me 03.09.2026 aty
qendronin 74 rreshta te "shansit te dyte" dhe tipa si "Softwarehersteller
mit IT-Service" ose "Cybersecurity-Dienstleister".

Cka ben:
  1. Merr cdo rresht te listave FERTIG te zonave te dhena.
  2. E gjen firmen ne dosjet e mbledhjes se zones dhe tekstin e faqes se
     saj te ruajtur (01-webtext.json) - asnje faqe nuk rikapet.
  3. E gjykon me pipeline.branchen_filter.firma_bewerten() - i njejti
     filter qe perdorin zonat 35-39 qe nga fillimi.
  4. Shkruan gjykimin e ri TE firmen.json i dosjes se mbledhjes
     (profil_passt/typ/grund, quelle="nachpruefung-<data>"), sepse aty
     eshte e verteta nga e cila lexojne lista dhe baza. Gjykimi i vjeter
     ruhet ne te njejtin rekord (profil_alt) - asgje nuk fshihet.
  5. Nje raport JSON te laeufe/ me cdo ndryshim dhe arsyen.

Perdorimi:
    python werkzeuge/listen-nachpruefung.py 32 33 34
    python werkzeuge/listen-nachpruefung.py 32 --prove    # vetem tregon

Kosto: nje thirrje KI per rresht (gpt-4.1-mini, rreth 0.001 USD).
"""
import glob
import json
import os
import re
import sys
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from pathlib import Path

PROJEKT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJEKT))

from pipeline.env import lade_dotenv  # noqa: E402

lade_dotenv(PROJEKT / ".env")

import openpyxl  # noqa: E402
from pipeline.branchen_filter import firma_bewerten  # noqa: E402
from pipeline.ki import KI  # noqa: E402

SOT = datetime.now().strftime("%Y-%m-%d")
QUELLE = f"nachpruefung-{SOT}"


def domain(w):
    o = re.sub(r"^https?://", "", str(w or "").strip().lower())
    return re.sub(r"^www\.", "", o).split("/")[0]


def log(*t):
    print(f"[{datetime.now():%H:%M:%S}]", *t, flush=True)


def lista_e_fundit(zona):
    fs = sorted(glob.glob(str(PROJEKT / f"IT-Liste-Emails-Zona{zona}-FERTIG-*.xlsx")),
                key=os.path.getmtime)
    return fs[-1] if fs else None


def dosjet_e_mbledhjes(zona):
    """Dosjet e zones me firmen.json, pa ato te Dropcontact-it."""
    return [d for d in sorted((PROJEKT / "laeufe/leadquellen").glob(f"zona{zona}-*"))
            if (d / "firmen.json").exists() and "-dropcontact-" not in d.name]


def main():
    zonat = [a for a in sys.argv[1:] if a.isdigit()]
    prove = "--prove" in sys.argv
    if not zonat:
        sys.exit("Duhet se paku nje zone, p.sh.: listen-nachpruefung.py 32 33 34")

    ki = KI()
    log(f"modeli: {ki.model}")
    raport = []

    for zona in zonat:
        lista = lista_e_fundit(zona)
        if not lista:
            log(f"zona {zona}: asnje liste FERTIG - kapercehet")
            continue
        ws = openpyxl.load_workbook(lista).active
        domains = {domain(ws.cell(r, 10).value) for r in range(2, ws.max_row + 1)}
        domains.discard("")

        # firma + teksti i faqes, nga dosja e pare qe e njeh firmen
        firmat, tekstet, ku = {}, {}, {}
        for d in dosjet_e_mbledhjes(zona):
            fs = json.loads((d / "firmen.json").read_text(encoding="utf-8"))
            wt = {}
            if (d / "01-webtext.json").exists():
                wt = json.loads((d / "01-webtext.json").read_text(encoding="utf-8"))
            for f in fs:
                k = domain(f.get("website"))
                if k in domains and k not in firmat:
                    firmat[k], ku[k] = f, d
                    tekstet[k] = wt.get(k, "") if isinstance(wt.get(k, ""), str) else ""

        mungojne = domains - set(firmat)
        log(f"zona {zona}: {len(domains)} firma ne liste, {len(firmat)} te gjetura, "
            f"{len(mungojne)} pa rekord")

        def gjyko(k):
            return k, firma_bewerten(firmat[k], tekstet.get(k, ""), ki)

        with ThreadPoolExecutor(max_workers=8) as ex:
            gjykime = dict(ex.map(gjyko, sorted(firmat)))

        ndryshuar = 0
        for k, g in gjykime.items():
            f = firmat[k]
            para = bool(f.get("profil_passt"))
            pas = bool(g.get("passt"))
            hyrja = {"zona": zona, "firma": f.get("name"), "domain": k,
                     "dosja": ku[k].name, "para": para, "pas": pas,
                     "typ_para": f.get("profil_typ", ""), "typ_pas": g.get("typ", ""),
                     "grund_pas": g.get("grund", "")}
            raport.append(hyrja)
            if para != pas:
                ndryshuar += 1
                log(f"  {'HIQET' if not pas else 'KTHEHET'}: {f.get('name')} "
                    f"| {g.get('typ')} | {g.get('grund', '')[:90]}")
        log(f"zona {zona}: {ndryshuar} gjykime te ndryshuara nga {len(gjykime)}")

        if prove:
            continue
        # shkruaj te firmen.json - vetem rekordet e gjykuara
        for d in {ku[k] for k in gjykime}:
            pfad = d / "firmen.json"
            fs = json.loads(pfad.read_text(encoding="utf-8"))
            for f in fs:
                k = domain(f.get("website"))
                if k in gjykime and ku[k] == d:
                    g = gjykime[k]
                    f["profil_alt"] = {"passt": f.get("profil_passt"),
                                       "typ": f.get("profil_typ"),
                                       "grund": f.get("profil_grund")}
                    f["profil_passt"] = bool(g.get("passt"))
                    f["profil_typ"] = g.get("typ", "")
                    f["profil_grund"] = g.get("grund", "")
                    f["profil_quelle"] = QUELLE
            pfad.write_text(json.dumps(fs, ensure_ascii=False, indent=1), encoding="utf-8")
            log(f"  u shkrua: {pfad.parent.name}/firmen.json")

    dalja = PROJEKT / "laeufe" / f"nachpruefung-{SOT}.json"
    dalja.write_text(json.dumps(raport, ensure_ascii=False, indent=1), encoding="utf-8")
    hiqen = [r for r in raport if r["para"] and not r["pas"]]
    log(f"RAPORTI: {dalja.name} | {len(raport)} gjykime | {len(hiqen)} dalin jashte"
        f"{' (PROVE - asgje s u shkrua)' if prove else ''}")


if __name__ == "__main__":
    main()
