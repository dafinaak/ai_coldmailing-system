"""Schritt 4 startet den echten Lauf - und nur den echten.

Der Assistent darf keine zweite, schwaecher gepruefte Strecke aufmachen.
Deshalb wird hier festgehalten, dass er dieselbe Pipeline startet wie
das alte Formular, ihr die ausgewaehlten Firmen mitgibt, und selbst
nichts versendet.
"""
import json

from fastapi.testclient import TestClient
from passlib.context import CryptContext
import pytest
import yaml

import web.routen.assistent as assistent_routen
from web.app import create_app

PWD_CONTEXT = CryptContext(schemes=["bcrypt"], deprecated="auto")


class FakeLaufmanager:
    """Merkt sich, womit der Lauf gestartet wurde - startet nichts."""
    letzter = None

    def __init__(self, daten_dir):
        self.daten_dir = daten_dir

    def starte(self, kunde_datei, limit, firmen_datei=None):
        FakeLaufmanager.letzter = {
            "kunde_datei": kunde_datei, "limit": limit,
            "firmen_datei": firmen_datei}
        lauf_dir = self.daten_dir / "laeufe" / "test-kunde" / "20260812-140000"
        lauf_dir.mkdir(parents=True, exist_ok=True)
        return lauf_dir


@pytest.fixture
def daten_dir(tmp_path, monkeypatch):
    monkeypatch.setenv("WEB_SECRET", "test-geheimnis-nur-fuer-tests")
    monkeypatch.setenv("WEB_COOKIE_SECURE", "0")
    (tmp_path / "users.yaml").write_text(yaml.safe_dump(
        [{"name": "Lena Hartmann", "passwort_hash": PWD_CONTEXT.hash("richtig123")}],
        allow_unicode=True), encoding="utf-8")
    quelle = tmp_path / "laeufe" / "leadquellen" / "testgebiet"
    quelle.mkdir(parents=True)
    (quelle / "firmen.json").write_text(json.dumps([
        {"name": "Nah IT GmbH", "plz": "30159", "domain": "nah.de",
         "website": "https://nah.de", "categories": ["IT-Berater"]},
        {"name": "Weit IT GmbH", "plz": "80331", "domain": "weit.de",
         "website": "https://weit.de", "categories": ["IT-Berater"]},
    ]), encoding="utf-8")
    return tmp_path


@pytest.fixture
def angemeldet(daten_dir, monkeypatch):
    FakeLaufmanager.letzter = None
    monkeypatch.setattr("web.laufmanager.Laufmanager", FakeLaufmanager)
    client = TestClient(create_app(daten_dir))
    client.post("/login", data={"name": "Lena Hartmann", "passwort": "richtig123"})
    client.cookies.set("intro_gesehen", "1")
    return client


def entwurf_bis_schritt_4(angemeldet):
    kennung = angemeldet.get("/assistent", follow_redirects=False
                             ).headers["location"].split("/")[2]
    angemeldet.post(f"/assistent/{kennung}/1", data={
        "name": "Test Kunde", "verkaeufer_url": "meine-firma.de",
        "absender_email": "ich@meine-firma.de", "anzahl_leads": "50",
        "anweisungen": "ruhig bleiben"}, follow_redirects=False)
    angemeldet.post(f"/assistent/{kennung}/2", data={
        "usp_titel_0": "Feste Ansprechpartner", "usp_text_0": "Dieselbe Person.",
        "icp_firmografisch": "Handwerk"}, follow_redirects=False)
    angemeldet.post(f"/assistent/{kennung}/3", data={
        "ort": "Hannover", "radius_km": "40", "dienste": "IT-Berater"},
        follow_redirects=False)
    return kennung


def test_schritt_4_startet_den_lauf_mit_den_gewaehlten_firmen(angemeldet, daten_dir):
    kennung = entwurf_bis_schritt_4(angemeldet)

    antwort = angemeldet.post(f"/assistent/{kennung}/4",
                              data={"firma": "nah.de"}, follow_redirects=False)

    assert antwort.status_code == 303
    assert FakeLaufmanager.letzter["limit"] == 1
    firmen = json.loads(
        (daten_dir / FakeLaufmanager.letzter["firmen_datei"]).read_text())
    assert [f["name"] for f in firmen] == ["Nah IT GmbH"]


def test_abgewaehlte_firmen_gehen_nicht_mit(angemeldet, daten_dir):
    kennung = entwurf_bis_schritt_4(angemeldet)

    angemeldet.post(f"/assistent/{kennung}/4",
                    data={"firma": ["nah.de", "weit.de"], "abgewaehlt": "weit.de"},
                    follow_redirects=False)

    firmen = json.loads(
        (daten_dir / FakeLaufmanager.letzter["firmen_datei"]).read_text())
    assert [f["name"] for f in firmen] == ["Nah IT GmbH"]


def test_ohne_auswahl_wird_nichts_gestartet(angemeldet):
    kennung = entwurf_bis_schritt_4(angemeldet)

    angemeldet.post(f"/assistent/{kennung}/4",
                    data={"firma": "nah.de", "abgewaehlt": "nah.de"},
                    follow_redirects=False)

    assert FakeLaufmanager.letzter is None


def test_kundendatei_traegt_die_antworten_des_assistenten(angemeldet, daten_dir):
    kennung = entwurf_bis_schritt_4(angemeldet)

    angemeldet.post(f"/assistent/{kennung}/4", data={"firma": "nah.de"},
                    follow_redirects=False)

    kunde = yaml.safe_load(
        (daten_dir / FakeLaufmanager.letzter["kunde_datei"]).read_text())
    assert kunde["name"] == "Test Kunde"
    assert "Feste Ansprechpartner" in kunde["angebot"]
    assert kunde["anbieter_reihenfolge"] == ["impressum"]
    # Sicherheitsnetz: bis jemand echte Testempfaenger eintraegt, darf
    # dieser Kunde nur an die eigene Adresse senden.
    assert kunde["test_empfaenger"] == ["ich@meine-firma.de"]


def test_schritt_6_zeigt_den_stand_und_den_weg_zur_freigabe(angemeldet):
    kennung = entwurf_bis_schritt_4(angemeldet)
    angemeldet.post(f"/assistent/{kennung}/4", data={"firma": "nah.de"},
                    follow_redirects=False)

    seite = angemeldet.get(f"/assistent/{kennung}/6").text

    assert "/auftraege/test-kunde/20260812-140000/status.json" in seite
    assert "/pruefen/test-kunde/20260812-140000" in seite
    assert "sendet nichts" in seite


def test_kundendatei_besteht_die_echte_pruefung(angemeldet, daten_dir):
    # Im echten Probelauf am 12.08.2026 brach der Lauf hier ab: Schritt 5
    # fragt ABSTAENDE (7 und 7), die Kundendatei will die Tage AB START
    # (7 und 14). Geschrieben wurde [7, 7] - und load_kunde lehnte das
    # zu Recht ab. Diese Pruefung laeuft jetzt schon im Test mit.
    from pipeline.config import load_kunde

    kennung = entwurf_bis_schritt_4(angemeldet)
    angemeldet.post(f"/assistent/{kennung}/5", data={
        "versand_postfach": "ich@meine-firma.de", "abstand_1_2": "7",
        "abstand_2_3": "7"}, follow_redirects=False)
    angemeldet.post(f"/assistent/{kennung}/4", data={"firma": "nah.de"},
                    follow_redirects=False)

    pfad = daten_dir / FakeLaufmanager.letzter["kunde_datei"]
    kunde = load_kunde(pfad)          # wirft, wenn die Datei ungueltig ist

    assert kunde.follow_up_tage == [7, 14]


def test_bestehende_kundendatei_wird_nicht_ueberschrieben(angemeldet, daten_dir):
    # "Test Kunde" im Assistenten darf eine gepflegte kunden/test-kunde.yaml
    # nicht zerschiessen - sonst waere ein Tippfehler im Kampagnennamen
    # genug, um eine echte Kundendatei zu verlieren.
    (daten_dir / "kunden").mkdir(exist_ok=True)
    alt = daten_dir / "kunden" / "test-kunde.yaml"
    alt.write_text("name: Wichtiger Bestandskunde\n", encoding="utf-8")

    kennung = entwurf_bis_schritt_4(angemeldet)
    angemeldet.post(f"/assistent/{kennung}/4", data={"firma": "nah.de"},
                    follow_redirects=False)

    assert "Wichtiger Bestandskunde" in alt.read_text()
    assert FakeLaufmanager.letzter["kunde_datei"] != "kunden/test-kunde.yaml"


def test_schritt_6_kennt_die_echten_zustandsnamen(angemeldet):
    # Erster Bau lauschte auf Felder, die es gar nicht gibt ("fertig",
    # "text") - der Knopf zur Freigabe waere nie erschienen. Die Namen
    # kommen woertlich aus Laufmanager.status().
    from web.laufmanager import Laufmanager

    kennung = entwurf_bis_schritt_4(angemeldet)
    angemeldet.post(f"/assistent/{kennung}/4", data={"firma": "nah.de"},
                    follow_redirects=False)

    seite = angemeldet.get(f"/assistent/{kennung}/6").text

    assert "wartet_auf_freigabe" in seite
    assert "laeuft" in seite
    assert "stand.fertig" not in seite


def test_excel_kommt_als_download_mit_den_kontakten(angemeldet, daten_dir):
    import io
    import json

    import openpyxl

    kennung = entwurf_bis_schritt_4(angemeldet)
    angemeldet.post(f"/assistent/{kennung}/4", data={"firma": "nah.de"},
                    follow_redirects=False)
    lauf = daten_dir / "laeufe" / "test-kunde" / "20260812-140000"
    (lauf / "leads.json").write_text(json.dumps({"leads": [
        {"first_name": "Tim", "last_name": "Cappelmann", "email": "t@nah.de",
         "company": "Nah IT GmbH", "title": "GF", "website": "https://nah.de",
         "source": "impressum", "notizen": []}]}), encoding="utf-8")
    (lauf / "firmen.json").write_text(json.dumps([
        {"name": "Nah IT GmbH", "domain": "nah.de", "website": "https://nah.de",
         "telefon": "0511 1", "plz": "30159", "ausgang": "mit_entscheider"}]),
        encoding="utf-8")

    antwort = angemeldet.get(f"/assistent/{kennung}/kontakte.xlsx")

    assert antwort.status_code == 200
    assert "attachment" in antwort.headers["content-disposition"]
    assert antwort.headers["content-disposition"].endswith('.xlsx"')
    wb = openpyxl.load_workbook(io.BytesIO(antwort.content))
    assert wb.sheetnames == ["Kontakte", "Anruf & Brief", "Zur Kontrolle"]
    assert wb["Kontakte"].cell(row=2, column=4).value == "Herr Cappelmann"


def test_excel_ohne_lauf_sagt_es_statt_leer_zu_liefern(angemeldet):
    kennung = angemeldet.get("/assistent", follow_redirects=False
                             ).headers["location"].split("/")[2]

    antwort = angemeldet.get(f"/assistent/{kennung}/kontakte.xlsx")

    assert antwort.status_code == 400
    assert "noch keine Suche" in antwort.text


def test_schritt_6_bietet_den_excel_download_an(angemeldet):
    kennung = entwurf_bis_schritt_4(angemeldet)
    angemeldet.post(f"/assistent/{kennung}/4", data={"firma": "nah.de"},
                    follow_redirects=False)

    seite = angemeldet.get(f"/assistent/{kennung}/6").text

    assert f"/assistent/{kennung}/kontakte.xlsx" in seite
    assert "Excel" in seite


def test_schritt_6_ohne_gestarteten_lauf_sagt_es(angemeldet):
    kennung = angemeldet.get("/assistent", follow_redirects=False
                             ).headers["location"].split("/")[2]

    seite = angemeldet.get(f"/assistent/{kennung}/6").text

    assert "noch keine Suche gestartet" in seite
