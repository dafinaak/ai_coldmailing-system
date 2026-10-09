"""All contacts of ten zones in one Excel sheet (Dafina, 29.09.2026).

The final list of every zone (IT-Liste-Emails-Zona<NR>-FERTIG-<time>.xlsx)
stays as it is. This puts the contacts of the newest list of each zone of a
range - 40-49, 50-59 ... - under each other on ONE sheet ("boni kejt te nje
sheet mos i ndaj hiq"). The columns stay exactly those of the zone lists, so
whatever reads them by position keeps working; one last column "Zona" says
where a row comes from, and "Nr" counts through the whole sheet.

Only what can be sent goes in. The catch-all addresses ("Bezahlt, nicht
versandfähig") and the check sheet ("Zur Kontrolle") stay on the zone lists:
an address nobody may send to must not sit between those that can be sent.
The check notes are on every row anyway, in the column "Hinweis".

Two things stop the merge instead of producing a file:
  - a zone whose list has other columns than the first zone's - its rows
    would land under the wrong headers;
  - the same person or e-mail on two zone lists - the zone lists already
    drop such a person (the lower zone keeps it), and if one slipped
    through, one campaign would write to them twice.

Reads only; nothing in the zone lists changes.
"""
from __future__ import annotations

import re
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

CONTACTS = "Versandfertig"
ZONE_LIST = re.compile(r"^IT-Liste-Emails-Zona(\d{1,2})-FERTIG-(\d{8}-\d{4})\.xlsx$")


def _in_dekade(zone: str, start: int) -> bool:
    """Gehoert diese Zone in die Dekade, die bei `start` beginnt?

    Ein Zonenname ist ein Postleitzahl-Anfang. Zweistellig ("83") heisst
    genau diese Zone; einstellig ("8") heisst die ganze Dekade - Muenchen
    bekam am 08.10.2026 den Namen "8", weil seine Codes ueber 80xxx,
    81xxx und 85xxx laufen. Ohne diesen Fall fehlte es stillschweigend in
    der Sammelliste 80-89, und zwar die groesste Zone des Projekts.
    """
    if len(zone) == 1:
        return str(start).startswith(zone)
    return start <= int(zone) <= start + 9


def newest_per_zone(folder, start: int) -> dict:
    """{"40": path, "41": path ...}: the newest final list of each zone from
    start to start + 9. Intermediate lists and merged files do not count."""
    newest = {}
    for path in Path(folder).glob("IT-Liste-Emails-Zona*-FERTIG-*.xlsx"):
        match = ZONE_LIST.match(path.name)
        if not match or not _in_dekade(match.group(1), start):
            continue
        zone, stamp = match.groups()
        if zone not in newest or stamp > newest[zone][0]:
            newest[zone] = (stamp, path)
    return {zone: path for zone, (_, path) in sorted(newest.items())}


def _header(sheet) -> list:
    return [cell.value for cell in sheet[1]]


def _check_people(sheets: dict) -> None:
    seen, twice = {}, []
    for zone, sheet in sheets.items():
        header = _header(sheet)
        for row in sheet.iter_rows(min_row=2, values_only=True):
            for column in ("E-Mail", "Person"):
                if column not in header:
                    continue
                value = str(row[header.index(column)] or "").strip().casefold()
                if not value:
                    continue
                if value in seen and seen[value] != zone:
                    twice.append(f"{value} (zona {seen[value]} dhe {zone})")
                seen.setdefault(value, zone)
    if twice:
        raise ValueError("I njejti person ne dy lista zonash: " + "; ".join(twice))


def _copy_style(source, target) -> None:
    target.font = copy(source.font)
    target.fill = copy(source.fill)
    target.border = copy(source.border)
    target.alignment = copy(source.alignment)
    target.number_format = source.number_format


def merge_zone_lists(files: dict, out_path) -> Path:
    """files: {zone: path of its final list}. Writes the merged workbook."""
    sheets = {}
    for zone, path in sorted(files.items()):
        book = openpyxl.load_workbook(path)
        if CONTACTS in book.sheetnames:
            sheets[zone] = book[CONTACTS]
    if not sheets:
        raise ValueError("Asnje liste zone me flete 'Versandfertig'.")
    first_zone, first = next(iter(sheets.items()))
    header = _header(first)
    for zone, sheet in sheets.items():
        if _header(sheet) != header:
            raise ValueError(f"Zona {zone}: lista ka kolona tjera se zona "
                             f"{first_zone} - bashkimi ndalet.")
    _check_people(sheets)

    merged = openpyxl.Workbook()
    target = merged.active
    target.title = CONTACTS
    target.append(header + ["Zona"])
    for column, cell in enumerate(first[1], 1):
        _copy_style(cell, target.cell(1, column))
        width = first.column_dimensions[get_column_letter(column)].width
        if width:
            target.column_dimensions[get_column_letter(column)].width = width
    _copy_style(first.cell(1, len(header)), target.cell(1, len(header) + 1))
    target.column_dimensions[get_column_letter(len(header) + 1)].width = 7
    target.row_dimensions[1].height = first.row_dimensions[1].height

    numbered = header[:1] == ["Nr"]
    nr = 0
    for zone, sheet in sheets.items():
        for row in sheet.iter_rows(min_row=2):
            if all(cell.value in (None, "") for cell in row):
                continue
            nr += 1
            values = [cell.value for cell in row]
            if numbered:
                values[0] = nr
            target.append(values + [zone])
            for column, cell in enumerate(row, 1):
                _copy_style(cell, target.cell(target.max_row, column))
    target.freeze_panes = "A2"
    target.auto_filter.ref = target.dimensions

    out_path = Path(out_path)
    merged.save(out_path)
    return out_path
