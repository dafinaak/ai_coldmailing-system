"""Die Excel-Tabelle aus einem fertigen Lauf.

Diese Datei geht aus der Hand - per Mail, in eine Ablage, zu Oliver.
Was hier fehlt, fehlt dort auch, ohne dass es jemand merkt. Deshalb
haengen die Tests vor allem an dem, was NICHT verschwinden darf: die
Firmen ohne Adresse und die Hinweise zur Kontrolle.
"""
import json

import openpyxl
import pytest

from pipeline.kontakte_excel import mappe_bauen, schreiben


def lauf_schreiben(ordner, leads, firmen, personalisierung=None):
    (ordner / "leads.json").write_text(json.dumps(
        {"leads": leads, "deckung": {}}), encoding="utf-8")
    (ordner / "firmen.json").write_text(json.dumps(firmen), encoding="utf-8")
    (ordner / "personalisierung.json").write_text(json.dumps(
        personalisierung or {"fertig": [], "nacharbeit": []}), encoding="utf-8")
    return ordner


def lead(vorname, nachname, mail, firma="A GmbH", website="https://a.de",
         notizen=()):
    return {"first_name": vorname, "last_name": nachname, "email": mail,
            "company": firma, "title": "Geschäftsführung", "website": website,
            "source": "impressum", "notizen": list(notizen)}


def firma(name="A GmbH", domain="a.de", ausgang="mit_entscheider", telefon="0511 1"):
    return {"name": name, "domain": domain, "website": f"https://{domain}",
            "plz": "30159", "telefon": telefon, "ausgang": ausgang,
            "stufe": "impressum"}


def blatt(wb, name):
    return [list(r) for r in wb[name].iter_rows(values_only=True)]


def test_kontakte_stehen_mit_anrede_und_telefon_drin(tmp_path):
    lauf_schreiben(tmp_path, [lead("Tim", "Cappelmann", "t@a.de")], [firma()])

    zeilen = blatt(mappe_bauen(tmp_path), "Kontakte")

    assert zeilen[0][:6] == ["Nr", "Firma", "Person", "Anrede", "E-Mail", "Telefon"]
    assert zeilen[1][1:6] == ["A GmbH", "Tim Cappelmann", "Herr Cappelmann",
                              "t@a.de", "0511 1"]


def test_sammeladresse_bekommt_die_anrede_ohne_namen(tmp_path):
    lauf_schreiben(tmp_path, [lead("", "", "info@a.de")], [firma()])

    zeilen = blatt(mappe_bauen(tmp_path), "Kontakte")

    assert zeilen[1][3] == "zusammen"
    assert "Sammeladresse" in zeilen[1][-1]


def test_firmen_ohne_adresse_landen_auf_anruf_und_brief(tmp_path):
    # Der wichtigste Fall: diese Firmen tauchen sonst nirgends auf und
    # sehen aus, als haette es sie nie gegeben.
    lauf_schreiben(tmp_path, [lead("Tim", "Cappelmann", "t@a.de")], [
        firma(),
        firma("B GmbH", "b.de", ausgang="info_ungueltig", telefon="0511 2"),
        firma("C GmbH", "c.de", ausgang="kein_entscheider", telefon="0511 3")])

    zeilen = blatt(mappe_bauen(tmp_path), "Anruf & Brief")

    assert [z[0] for z in zeilen[1:]] == ["B GmbH", "C GmbH"]
    assert "nicht zustellbar" in zeilen[1][1]
    assert zeilen[1][2] == "0511 2"          # Telefonnummer muss mit


def test_erreichte_firmen_stehen_nicht_auf_anruf_und_brief(tmp_path):
    lauf_schreiben(tmp_path, [lead("Tim", "Cappelmann", "t@a.de")],
                   [firma(), firma("B GmbH", "b.de", ausgang="info_fallback")])

    zeilen = blatt(mappe_bauen(tmp_path), "Anruf & Brief")

    assert len(zeilen) == 1                  # nur die Kopfzeile


def test_abweichende_maildomain_wird_vermerkt(tmp_path):
    lauf_schreiben(tmp_path, [lead("Tim", "Cappelmann", "t@fremd.de")], [firma()])

    zeilen = blatt(mappe_bauen(tmp_path), "Kontakte")

    assert "fremd.de" in zeilen[1][-1]
    assert len(blatt(mappe_bauen(tmp_path), "Zur Kontrolle")) == 2


def test_notizen_aus_dem_lauf_gehen_nicht_verloren(tmp_path):
    lauf_schreiben(tmp_path, [lead("Tim", "Cappelmann", "t@a.de",
                                   notizen=["Dropcontact führt ihn anders"])],
                   [firma()])

    assert "anders" in blatt(mappe_bauen(tmp_path), "Kontakte")[1][-1]


def test_textstand_wird_gezeigt(tmp_path):
    lauf_schreiben(
        tmp_path,
        [lead("Tim", "Cappelmann", "t@a.de"), lead("Ann", "Beck", "b@a.de")],
        [firma()],
        {"fertig": [{"email": "t@a.de", "betreff": "Kurz gefragt"}],
         "nacharbeit": [{"email": "b@a.de", "betreff": "Hallo",
                         "grund": "zu lang"}]})

    zeilen = blatt(mappe_bauen(tmp_path), "Kontakte")

    assert zeilen[1][8:10] == ["Kurz gefragt", "fertig"]
    assert zeilen[2][9] == "Nacharbeit nötig"
    assert "zu lang" in zeilen[2][-1]


def test_leerer_lauf_ergibt_eine_leere_aber_gueltige_datei(tmp_path):
    lauf_schreiben(tmp_path, [], [])

    wb = mappe_bauen(tmp_path)

    assert wb.sheetnames == ["Kontakte", "Anruf & Brief", "Zur Kontrolle"]
    assert len(blatt(wb, "Kontakte")) == 1


def test_fehlende_dateien_stuerzen_nicht_ab(tmp_path):
    wb = mappe_bauen(tmp_path)      # gar kein Laufordner-Inhalt

    assert len(blatt(wb, "Kontakte")) == 1


def test_schreiben_legt_die_datei_an(tmp_path):
    lauf_schreiben(tmp_path, [lead("Tim", "Cappelmann", "t@a.de")], [firma()])
    ziel = tmp_path / "kontakte.xlsx"

    schreiben(tmp_path, ziel)

    assert openpyxl.load_workbook(ziel)["Kontakte"].max_row == 2
