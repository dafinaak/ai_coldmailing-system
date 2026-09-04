#!/usr/bin/env python3
"""Kontrolli i listave FERTIG para se t'i shkojne klientit.

Porosi e Dafines, 03.09.2026: "boni check websites, emailat, numrat e
tel ... kqyri kejt sakte se kom me ja dergu oliverit".

Cka kontrollon, per cdo rresht te cdo liste:
  - faqja:      e lexueshme si adrese, me domain; dhe A HAPET vertet
                (kerkese HTTP e gjalle, pa kosto)
  - email-i:    format i vlefshem; jo adrese gjenerale; domain-i i tij
                perputhet me domain-in e faqes (nje email personal ne
                domain tjeter eshte i dyshimte)
  - telefonat:  7-15 shifra pas pastrimit - as te cunguar, as te gjate
  - personi:    emer + mbiemer; Anrede fillon me Herr/Frau
  - kodi:       5 shifra, brenda zones
  - dublikata:  i njejti email ose i njejti njeri ne dy lista

Email-at NUK riverifikohen ketu: erdhen te verifikuar nga Dropcontact
("gebaut + verifiziert") dhe cdo riverifikim do te kushtonte nje kredit
per adrese. Faqet po - ato kontrollohen te gjalla, sepse nje domain i
vdekur do te thote qe firma ndoshta s'ekziston me.

Dalja: nje Excel "Listen-Pruefbericht-<koha>.xlsx" me nje rresht per
cdo gjetje, plus permbledhje ne ekran. Asgje nuk ndryshohet ne lista -
kjo vegel vetem lexon.
"""
import glob
import os
import re
import sys
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from pathlib import Path

PROJEKT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJEKT))

import openpyxl  # noqa: E402
import requests  # noqa: E402

GJENERALE = re.compile(
    r"^(info|kontakt|office|mail|service|hallo|team|zentrale|anfrage|"
    r"vertrieb|support|buero|kontact|post|verwaltung)@", re.I)
EMAIL = re.compile(r"^[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}$")
KOKA_UA = {"User-Agent": "Mozilla/5.0 (compatible; Listenpruefung/1.0)"}


def domain(w):
    o = re.sub(r"^https?://", "", str(w or "").strip().lower())
    return re.sub(r"^www\.", "", o).split("/")[0]


def shifra(t):
    d = re.sub(r"\D", "", str(t or ""))
    if d.startswith("00"):
        d = d[2:]
    if d.startswith("49"):
        d = d[2:]
    return d.lstrip("0")


def faqja_hapet(url):
    """(status, shenim). Provohet HEAD, pastaj GET; https, pastaj http."""
    d = domain(url)
    if not d:
        return "PA DOMAIN", ""
    for skema in ("https://", "http://"):
        for metoda in ("head", "get"):
            try:
                r = getattr(requests, metoda)(skema + d, headers=KOKA_UA,
                                              timeout=8, allow_redirects=True)
                if r.status_code < 400:
                    return "OK", f"{r.status_code}"
                if r.status_code in (403, 405, 429) and metoda == "head":
                    continue           # disa servera e refuzojne HEAD-in
                if r.status_code >= 500 or r.status_code == 404:
                    return "PROBLEM", f"HTTP {r.status_code}"
                return "OK", f"{r.status_code}"
            except requests.exceptions.SSLError:
                break                  # provo http
            except requests.RequestException as g:
                gabim = type(g).__name__
        # kalon te skema tjeter
    return "NUK HAPET", gabim if "gabim" in dir() else "pa pergjigje"


def listat_e_fundit():
    fs = {}
    for f in glob.glob(str(PROJEKT / "IT-Liste-Emails-Zona3*-FERTIG-*.xlsx")):
        z = os.path.basename(f).split("-")[3][-2:]
        if z not in fs or os.path.getmtime(f) > os.path.getmtime(fs[z]):
            fs[z] = f
    return dict(sorted(fs.items()))


def main():
    gjetje = []          # (zona, nr, firma, fusha, problemi, vlera)
    email_ku = {}        # email -> (zona, nr, firma)
    person_ku = {}       # (person, domain) -> (zona, nr)
    faqet = {}           # domain -> url (per kontrollin e gjalle)
    rreshta = 0

    for z, f in listat_e_fundit().items():
        ws = openpyxl.load_workbook(f).active
        for r in range(2, ws.max_row + 1):
            rreshta += 1
            nr, firma = ws.cell(r, 1).value, str(ws.cell(r, 2).value or "")
            person = str(ws.cell(r, 3).value or "").strip()
            email = str(ws.cell(r, 5).value or "").strip()
            anrede = str(ws.cell(r, 6).value or "")
            tel_p, tel_f = ws.cell(r, 8).value, ws.cell(r, 9).value
            web = str(ws.cell(r, 10).value or "").strip()
            plz = str(ws.cell(r, 11).value or "").strip()

            def g(fusha, problemi, vlera=""):
                gjetje.append((z, nr, firma, fusha, problemi, str(vlera)))

            if not firma:
                g("Firma", "bosh")
            if len(person.split()) < 2:
                g("Person", "pa emer e mbiemer te plote", person)
            if not (anrede.startswith("Herr ") or anrede.startswith("Frau ")):
                g("Anrede", "nuk fillon me Herr/Frau", anrede)

            if not EMAIL.match(email):
                g("E-Mail", "format i pavlefshem", email)
            elif GJENERALE.match(email):
                g("E-Mail", "adrese gjenerale, jo personale", email)
            else:
                de, dw = email.split("@")[1].lower(), domain(web)
                if dw and not (de == dw or de.endswith("." + dw) or dw.endswith("." + de)):
                    g("E-Mail", "domain-i i email-it s'perputhet me faqen",
                      f"{email}  <>  {dw}")
                if email.lower() in email_ku:
                    g("E-Mail", "DUBLIKATE - i njejti email edhe te",
                      "zona %s nr %s (%s)" % email_ku[email.lower()])
                else:
                    email_ku[email.lower()] = (z, nr, firma)

            for etiketa, tel in (("Telefon (Person)", tel_p), ("Telefon (Firma)", tel_f)):
                if tel:
                    n = len(shifra(tel))
                    if n < 7:
                        g(etiketa, "numer i cunguar (< 7 shifra)", tel)
                    elif n > 15:
                        g(etiketa, "numer shume i gjate (> 15 shifra)", tel)
            if not tel_p and not tel_f:
                g("Telefon", "asnje numer")

            if not domain(web):
                g("Webseite", "bosh ose e palexueshme", web)
            else:
                faqet.setdefault(domain(web), web)

            if not (plz.isdigit() and len(plz) == 5):
                g("PLZ", "jo 5 shifra", plz)
            elif not plz.startswith(z):
                g("PLZ", f"jashte zones {z}", plz)

            kyc = (person.casefold(), domain(web))
            if person and kyc in person_ku:
                g("Person", "DUBLIKATE - i njejti njeri edhe te",
                  "zona %s nr %s" % person_ku[kyc])
            else:
                person_ku[kyc] = (z, nr)

    print(f"Rreshta te kontrolluar: {rreshta} | faqe unike: {len(faqet)}")
    print("Po hapen faqet (kerkese e gjalle, pa kosto) ...")
    with ThreadPoolExecutor(max_workers=16) as ex:
        rez = dict(zip(faqet, ex.map(faqja_hapet, faqet.values())))
    keq = {d: r for d, r in rez.items() if r[0] != "OK"}
    # gjetjet e faqeve shkruhen per cdo rresht qe e perdor ate domain
    for z, f in listat_e_fundit().items():
        ws = openpyxl.load_workbook(f).active
        for r in range(2, ws.max_row + 1):
            d = domain(ws.cell(r, 10).value)
            if d in keq:
                gjetje.append((z, ws.cell(r, 1).value, str(ws.cell(r, 2).value or ""),
                               "Webseite", f"{keq[d][0]}: {keq[d][1]}", ws.cell(r, 10).value))

    # ------------------------------------------------ raporti
    wb = openpyxl.Workbook()
    b = wb.active
    b.title = "Gjetjet"
    b.append(["Zona", "Nr", "Firma", "Fusha", "Problemi", "Vlera"])
    for gj in sorted(gjetje, key=lambda x: (x[0], x[1] or 0)):
        b.append(list(gj))
    for c, w in zip("ABCDEF", (6, 5, 36, 18, 44, 44)):
        b.column_dimensions[c].width = w
    b.freeze_panes = "A2"
    b.auto_filter.ref = b.dimensions
    dalja = PROJEKT / f"Listen-Pruefbericht-{datetime.now():%Y%m%d-%H%M}.xlsx"
    wb.save(dalja)

    from collections import Counter
    print("\nGJETJET sipas fushes:")
    for k, v in Counter((gj[3], gj[4].split(":")[0]) for gj in gjetje).most_common():
        print(f"  {v:>4}  {k[0]:<18} {k[1]}")
    print(f"\nGjithsej gjetje: {len(gjetje)} ne {rreshta} rreshta")
    print(f"RAPORTI: {dalja.name}")


if __name__ == "__main__":
    main()
